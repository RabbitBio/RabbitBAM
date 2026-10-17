#!/usr/bin/env python3
"""Editable method figures inspired by the visual grammar of reference/plot.pptx.

Only the PPTX is published beside this script. SVG/PNG previews go to /tmp.
All objects are native PowerPoint text, shapes, or connectors; no raster artwork.
Run with python-pptx, Pillow and cairosvg installed. Re-running overwrites the
generated deck, so preserve any subsequent manual PowerPoint edits separately.
"""

from html import escape
from pathlib import Path
import math
import subprocess

import cairosvg
from PIL import ImageFont
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, MSO_AUTO_SIZE, PP_ALIGN
from pptx.enum.dml import MSO_LINE_DASH_STYLE
from pptx.oxml.xmlchemy import OxmlElement
from pptx.util import Inches, Pt


OUT = Path(__file__).resolve().parent
PREVIEW = Path('/tmp/swbam-redesign/final')
W, H = 1600, 900
INK = '173948'
TEXT = '172C36'
GRAY = '61727A'
LIGHT = 'A9B8BD'
BLUE = 'D5EDF8'
BLUE_DARK = '197BA3'
GREEN = 'DFEED2'
GREEN_DARK = '57834A'
YELLOW = 'FFF7CE'
GOLD = '9A7625'
ROSE = 'F7DEE1'
RED = 'B95362'
PURPLE = 'E8E1F2'
VIOLET = '7A6198'
PALE = 'F5F8F9'
WHITE = 'FFFFFF'


class Canvas:
    def __init__(self, title, notes):
        self.title, self.notes, self.items = title, notes, []

    def shape(self, name, x, y, w, h, fill=WHITE, stroke=INK,
              kind='rect', width=1.6, dash=False, radius=8):
        self.items.append(dict(type='shape', name=name, x=x, y=y, w=w, h=h,
                               fill=fill, stroke=stroke, kind=kind, width=width,
                               dash=dash, radius=radius))

    def text(self, name, x, y, w, h, s, size=21, bold=False, color=TEXT,
             align='center', italic=False):
        self.items.append(dict(type='text', name=name, x=x, y=y, w=w, h=h,
                               text=s, size=size, bold=bold, color=color,
                               align=align, italic=italic))

    def line(self, name, pts, color=INK, width=1.8, head=False, dash=False):
        self.items.append(dict(type='line', name=name, pts=pts, color=color,
                               width=width, head=head, dash=dash))

    def arrow(self, name, x1, y1, x2, y2, color=INK, dash=False, width=1.8):
        self.line(name, [(x1,y1),(x2,y2)], color, width, True, dash)

    def box(self, name, x, y, w, h, s, fill=WHITE, size=21, bold=False,
            stroke=INK, rounded=False):
        self.shape(name, x,y,w,h,fill,stroke,kind='round' if rounded else 'rect')
        self.text(name+' / label', x+6,y+3,w-12,h-6,s,size,bold)

    def panel(self, name, x,y,w,h, title, fill=PALE, dash=False):
        self.shape(name,x,y,w,h,fill,INK,kind='round',dash=dash,radius=12)
        self.text(name+' / heading',x+15,y+10,w-30,30,title,24,True,align='left')

    def section(self, y, s):
        self.text(s,30,y,1538,32,s,26,True,align='left')

    def chips(self, name,x,y,values,colors=None,cw=48,ch=35,gap=0,size=20):
        for i,v in enumerate(values):
            col=colors[i] if isinstance(colors,list) else colors or BLUE
            self.box(f'{name} / {i}',x+i*(cw+gap),y,cw,ch,str(v),col,size)

    def blocks(self,name,x,y,n,cw=36,ch=30,color=BLUE,selected=(),gap=3):
        for i in range(n):
            col=GREEN if i in selected else color
            self.shape(f'{name} / {i}',x+i*(cw+gap),y,cw,ch,col,INK,width=1.5)
            self.line(f'{name} / divider {i}',[(x+i*(cw+gap)+cw*.27,y),
                       (x+i*(cw+gap)+cw*.27,y+ch)],INK,1)

    def bracket(self,name,x1,x2,y,label,color=GRAY,below=True):
        dy=-7 if below else 7
        self.line(name,[(x1,y+dy),(x1,y),(x2,y),(x2,y+dy)],color,1.4)
        self.text(name+' / label',x1,y+6 if below else y-28,x2-x1,23,
                  label,18,color=color)

    def document(self,name,x,y,w,h,label,color=BLUE):
        self.shape(name,x,y,w,h,color,INK,kind='paper')
        self.text(name+' / label',x+3,y+h*.32,w-6,h*.45,label,19,True)

    def validate(self):
        font_cache={}
        for item in self.items:
            if item['type']=='line':
                assert all(0<=x<=W and 0<=y<=H for x,y in item['pts']),item['name']
                continue
            x,y,w,h=(item[k] for k in ('x','y','w','h'))
            assert x>=0 and y>=0 and w>0 and h>0 and x+w<=W and y+h<=H,item['name']
            if item['type']!='text': continue
            key=(item['size'],item['bold'],item['italic'])
            if key not in font_cache:
                style=('Bold' if item['bold'] else '')+(' Italic' if item['italic'] else '')
                path=subprocess.check_output(['fc-match','-f','%{file}',
                      'Liberation Sans:style='+style.strip()],text=True)
                font_cache[key]=ImageFont.truetype(path,round(item['size']*4))
            font=font_cache[key]
            for line in item['text'].split('\n'):
                length=font.getlength(line)/4
                assert length <= w-1, (item['name'],line,round(length,1),w)
            assert len(item['text'].split('\n'))*item['size']*1.12<=h+2,item['name']

    def svg(self):
        out=[f'<svg xmlns="http://www.w3.org/2000/svg" width="{W}" height="{H}" viewBox="0 0 {W} {H}">',
             f'<title>{escape(self.title)}</title><rect width="{W}" height="{H}" fill="white"/>']
        for it in self.items:
            out.append(f'<g><title>{escape(it["name"])}</title>')
            if it['type']=='shape':
                x,y,w,h=(it[k] for k in ('x','y','w','h'))
                style=f'fill="#{it["fill"]}" stroke="#{it["stroke"]}" stroke-width="{it["width"]}"'
                if it['dash']: style+=' stroke-dasharray="7 5"'
                if it['kind']=='ellipse':
                    out.append(f'<ellipse cx="{x+w/2}" cy="{y+h/2}" rx="{w/2}" ry="{h/2}" {style}/>')
                elif it['kind']=='paper':
                    f=min(w,h)*.23
                    out.append(f'<path d="M{x},{y} H{x+w} V{y+h-f} L{x+w-f},{y+h} H{x} Z" {style}/>')
                    out.append(f'<path d="M{x+w},{y+h-f} H{x+w-f} V{y+h}" fill="none" stroke="#{it["stroke"]}" stroke-width="{it["width"]}"/>')
                elif it['kind']=='diamond':
                    pts=[(x+w/2,y),(x+w,y+h/2),(x+w/2,y+h),(x,y+h/2)]
                    out.append(f'<polygon points="'+ ' '.join(f'{xx},{yy}' for xx,yy in pts)+f'" {style}/>')
                else:
                    r=it['radius'] if it['kind']=='round' else 0
                    out.append(f'<rect x="{x}" y="{y}" width="{w}" height="{h}" rx="{r}" {style}/>')
            elif it['type']=='text':
                lines=it['text'].split('\n'); size=it['size']; lh=size*1.12
                yy=it['y']+it['h']/2-(len(lines)-1)*lh/2+size*.35
                xx=it['x']+{'left':0,'center':it['w']/2,'right':it['w']}[it['align']]
                anchor={'left':'start','center':'middle','right':'end'}[it['align']]
                for s in lines:
                    out.append(f'<text x="{xx}" y="{yy}" font-family="Arial,Liberation Sans,sans-serif" font-size="{size}" font-weight="{"bold" if it["bold"] else "normal"}" font-style="{"italic" if it["italic"] else "normal"}" text-anchor="{anchor}" fill="#{it["color"]}">{escape(s)}</text>')
                    yy+=lh
            else:
                pts=it['pts'];dash=' stroke-dasharray="6 4"' if it['dash'] else ''
                out.append('<polyline points="'+' '.join(f'{x},{y}' for x,y in pts)+f'" fill="none" stroke="#{it["color"]}" stroke-width="{it["width"]}" stroke-linejoin="round"{dash}/>')
                if it['head']:
                    (x0,y0),(x,y)=pts[-2:];a=math.atan2(y-y0,x-x0);l=9;half=4
                    bx,by=x-l*math.cos(a),y-l*math.sin(a)
                    tip=[(x,y),(bx-half*math.sin(a),by+half*math.cos(a)),(bx+half*math.sin(a),by-half*math.cos(a))]
                    out.append('<polygon points="'+' '.join(f'{xx},{yy}' for xx,yy in tip)+f'" fill="#{it["color"]}"/>')
            out.append('</g>')
        return '\n'.join(out+['</svg>'])

    def add_slide(self,prs):
        slide=prs.slides.add_slide(prs.slide_layouts[6])
        slide.background.fill.solid();slide.background.fill.fore_color.rgb=RGBColor.from_string(WHITE)
        slide.notes_slide.notes_text_frame.text=self.title+'\n\n'+self.notes
        u=lambda x: Inches(x/120)
        for it in self.items:
            if it['type']=='text':
                sh=slide.shapes.add_textbox(u(it['x']),u(it['y']),u(it['w']),u(it['h']))
                tf=sh.text_frame;tf.clear();tf.word_wrap=False;tf.auto_size=MSO_AUTO_SIZE.NONE
                tf.margin_left=tf.margin_right=tf.margin_top=tf.margin_bottom=0
                tf.vertical_anchor=MSO_ANCHOR.MIDDLE
                for i,s in enumerate(it['text'].split('\n')):
                    pa=tf.paragraphs[0] if not i else tf.add_paragraph()
                    pa.alignment={'left':PP_ALIGN.LEFT,'center':PP_ALIGN.CENTER,'right':PP_ALIGN.RIGHT}[it['align']]
                    pa.space_before=pa.space_after=Pt(0);pa.line_spacing=Pt(it['size']*.6*1.12)
                    r=pa.add_run();r.text=s;r.font.name='Arial';r.font.size=Pt(it['size']*.6)
                    r.font.bold=it['bold'];r.font.italic=it['italic'];r.font.color.rgb=RGBColor.from_string(it['color'])
                sh.name=it['name']
                # Keep a block and its label movable as one editable object.
                if it['name'].endswith(' / label') and len(slide.shapes)>1:
                    prev=slide.shapes[-2]
                    if prev.name==it['name'][:-8]:
                        group=slide.shapes.add_group_shape([prev,sh])
                        group.name=prev.name+' [editable group]'
            elif it['type']=='shape':
                kind={'rect':MSO_SHAPE.RECTANGLE,'round':MSO_SHAPE.ROUNDED_RECTANGLE,
                      'ellipse':MSO_SHAPE.OVAL,'paper':MSO_SHAPE.FOLDED_CORNER,
                      'diamond':MSO_SHAPE.DIAMOND}[it['kind']]
                sh=slide.shapes.add_shape(kind,u(it['x']),u(it['y']),u(it['w']),u(it['h']))
                if it['kind']=='round': sh.adjustments[0]=min(.5,it['radius']/min(it['w'],it['h']))
                if it['kind']=='paper' and len(sh.adjustments): sh.adjustments[0]=.23
                sh.name=it['name'];sh.fill.solid();sh.fill.fore_color.rgb=RGBColor.from_string(it['fill'])
                sh.line.color.rgb=RGBColor.from_string(it['stroke']);sh.line.width=Pt(it['width']*.6)
                if not it['width']: sh.line.fill.background()
                if it['dash']: sh.line.dash_style=MSO_LINE_DASH_STYLE.DASH
                sh._element.spPr.append(OxmlElement('a:effectLst'))
            else:
                for i,(p,q) in enumerate(zip(it['pts'],it['pts'][1:])):
                    sh=slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT,u(p[0]),u(p[1]),u(q[0]),u(q[1]))
                    sh.name=it['name']+f' / segment {i}'
                    sh.line.color.rgb=RGBColor.from_string(it['color']);sh.line.width=Pt(it['width']*.6)
                    sh._element.spPr.append(OxmlElement('a:effectLst'))
                    if it['dash']: sh.line.dash_style=MSO_LINE_DASH_STYLE.DASH
                    if it['head'] and i==len(it['pts'])-2:
                        el=OxmlElement('a:tailEnd');el.set('type','triangle');el.set('w','med');el.set('len','med')
                        sh._element.spPr.get_or_add_ln().append(el)
        return slide


def overview():
    c=Canvas('SWBAM: library architecture and batch-level composition',
        'Redrawn using plot.pptx visual vocabulary: memory plane, MPE/CPE array, expanded processing components. '
        'Shapes illustrate SWBAM, not ZFP; no reference-paper mechanisms are claimed. '
        'One MPI rank uses one active core group (1 MPE + 64 CPEs). Four alternative processing pipelines, '
        'not four concurrent kernels. Application composition occurs at batch boundaries on the MPE. '
        'The common BAM read operator can emit raw views, counters, parsed records, or algorithm metadata. '
        'The optional bam1_t adapter materializes objects on the MPE and is not shown as a CPE stage. '
        'The BAM write pipeline also accepts already packed raw payloads (compression-only). '
        'Backend labels are for BAM; SAM backend support is not implied by the entire list. '
        'Source: swbam-sdk/lib/cpe/swbam_cpe_bam_read_pipeline.cpp, swbam_cpe_bam_write_pipeline.cpp; '
        'lib/operators/swbam_bam_transform_operator.cpp; include/swbam/cpe_pipeline.h; '
        'overleaf/SWBAM/bare_jrnl_new_sample4.tex Methods.')
    c.section(12,'(a) SWBAM library: data access and batch execution')
    c.shape('Data access plane',28,57,1544,138,BLUE,INK,width=2)
    c.text('Data access title',48,77,190,29,'Data access',25,True)
    c.text('Input backends',48,115,201,56,'BAM backends\nMemory / POSIX / MPI-IO',17)
    c.document('BAM input file',275,83,64,82,'BAM')
    c.arrow('Input to block batches',346,122,383,122)
    c.blocks('BGZF input blocks',393,95,6,cw=37,ch=43)
    c.bracket('Rank span',393,630,149,'Contiguous rank span')
    c.arrow('Blocks to alternating input',643,122,685,122)
    c.text('Input batch title',703,70,254,25,'Alternating batch buffers',19,True)
    c.box('Input A',702,106,111,47,'A',YELLOW,23,True)
    c.box('Input B',829,106,111,47,'B',GREEN,23,True)
    c.text('Batch capacity',702,159,238,24,'Up to 64 BGZF blocks',17)
    c.line('Separate input and output',[(980,77),(980,176)],LIGHT,1.4)
    c.text('Body output title',1004,71,264,25,'Rank body sink',20,True)
    c.blocks('Rank output blocks',1018,108,5,cw=38,ch=36,color=ROSE)
    c.text('Body strategies',1005,153,260,24,'Memory / spool',18)
    c.arrow('Sink to output',1276,124,1315,124)
    c.document('Final BAM file',1328,89,67,77,'BAM',ROSE)
    c.text('Distributed output label',1410,94,141,67,'Distributed\noutput\nMemory / MPI-IO',17)

    c.panel('Core group boundary',28,232,1544,345,'One MPI rank: 1 MPE + 64 CPEs',WHITE,True)
    c.panel('MPE controller',48,283,247,268,'MPE',YELLOW)
    for y,s in [(337,'Prepare batch'),(397,'Launch / join'),(457,'Batch post-processing')]:
        c.box('MPE '+s,65,y,213,44,s,WHITE,20,stroke=GOLD,rounded=True)
    c.text('Control granularity',66,514,211,24,'Once per batch',18,color=GOLD)
    c.arrow('MPE launches core group',295,405,329,405)
    c.text('CPE array title',337,283,207,28,'8 × 8 CPE array',21,True)
    for row in range(8):
        for col in range(8):
            c.shape(f'CPE {row*8+col}',345+col*24,326+row*24,21,21,
                    BLUE_DARK if (row,col)==(3,5) else GREEN,INK,width=1)
    c.text('CPE array range',334,526,215,24,'CPE₀ … CPE₆₃',18)
    c.line('Array expansion top',[(536,326),(581,284)],GRAY,1.3,dash=True)
    c.line('Array expansion bottom',[(536,515),(581,550)],GRAY,1.3,dash=True)
    c.panel('Reusable CPE components',582,283,969,268,'CPE processing components',GREEN)
    c.text('Kernel selection note',1112,294,421,26,'Selected per pipeline / operator',18,color=GREEN_DARK,align='right')
    # Four distinct pipeline options; never a simultaneous four-kernel schedule.
    rows=[(603,340,'BAM read',[('Decode',BLUE),('Parse',BLUE),('Record op.',YELLOW)]),
          (1080,340,'BAM write',[('Serialize',ROSE),('Compress',ROSE)]),
          (603,448,'SAM read',[('Text chunks',BLUE),('Parse',BLUE)]),
          (1080,448,'SAM write',[('BAM records',ROSE),('Format',ROSE)])]
    for x,y,name,stages in rows:
        c.text(name+' heading',x,y-21,438,26,name,20,True,align='left')
        cell=126 if len(stages)==3 else 192
        for i,(s,col) in enumerate(stages):
            xx=x+i*(cell+24)
            c.box(name+' '+s,xx,y+17,cell,47,s,col,19,stroke=INK,rounded=True)
            if i<len(stages)-1: c.arrow(name+f' stage {i}',xx+cell,y+40,xx+cell+24,y+40)
    c.arrow('Read batches into core group',757,195,757,279,BLUE_DARK,width=2.8)
    c.text('Read batch transfer label',774,202,221,24,'Batch input',18,color=BLUE_DARK,align='left')
    c.arrow('Compressed batches to sink',1177,279,1177,195,RED,width=2.8)
    c.text('Write batch transfer label',1193,202,243,24,'Batch output',18,color=RED,align='left')

    c.section(601,'(b) Compose applications at batch boundaries')
    chains=[(45,651,'Record counting',[('BAM read',BLUE),('Aggregate',YELLOW),('MPI reduce',PURPLE)]),
            (841,651,'BAM filtering / rewriting',[('BAM read + filter',BLUE),('Packing plan',YELLOW),('BAM write',ROSE)]),
            (45,753,'BAM → SAM',[('BAM read',BLUE),('Batch handoff',YELLOW),('SAM write',ROSE)]),
            (841,753,'SAM → BAM',[('SAM read',BLUE),('Packing plan',YELLOW),('BAM write',ROSE)])]
    for x,y,s,parts in chains:
        c.text(s+' title',x,y,702,26,s,21,True,align='left')
        for i,(label,col) in enumerate(parts):
            xx=x+i*237
            c.box(s+' / '+label,xx,y+33,207,43,label,col,19,rounded=True)
            if i<2: c.arrow(s+f' / link {i}',xx+207,y+54,xx+237,y+54)
    c.text('Higher level state',37,859,1526,27,
           'Sort / collate / fixmate / markdup add ordering, grouping, cross-record state and MPI exchange.',20,color=GRAY)
    return c


def timing():
    c=Canvas('Alternating buffers and MPE-CPE overlap',
        'Independent read and write schedules for one active core group; horizontal lengths are schematic, '
        'not measured times. Read shows overlap_input=true and prefetch_during_post_process=false. '
        'The MPE prefetches next input while the CPE handles current input. Post-processing follows join '
        'and validation; it is not overlapped with that read kernel. Write Start launches current compression '
        'then FlushPending writes previous output, optional prefetch follows, Complete joins and validates. '
        'No simultaneous read/write CPE kernels are implied. Raw payloads skip serialization. '
        'Source: swbam-sdk/lib/cpe/swbam_cpe_bam_read_pipeline.cpp RunCpeReadPipeline; '
        'swbam-sdk/lib/cpe/swbam_cpe_bam_write_pipeline.cpp CpeWritePipelineSession::{Start,Complete,FlushPending}.')
    for base,name in [(0,'BAM read'),(444,'BAM write')]:
        y=base
        c.section(y+12,f'({"a" if base==0 else "b"}) {name}: reuse buffers, overlap independent work')
        c.panel(name+' buffers',30,y+73,260,292,'Alternating buffers',BLUE if base==0 else ROSE)
        c.box(name+' buffer A',57,y+134,205,51,'A: batch i',YELLOW,23,True)
        c.box(name+' buffer B',57,y+212,205,51,
              'B: batch i+1' if base==0 else 'B: batch i−1',GREEN,23,True)
        c.text(name+' role',46,y+285,229,51,
               'Compressed input\nA / B swap after each batch' if base==0 else 'Compressed output\nCurrent / pending alternate',18)
        c.arrow(name+' time axis',393,y+107,1547,y+107,GRAY)
        c.text(name+' time label',1298,y+70,243,25,'Time (schematic)',18,color=GRAY,align='right')
        c.text(name+' MPE label',305,y+166,87,32,'MPE',23,True)
        c.text(name+' CPE label',298,y+281,96,32,'CPEs',23,True)
        # Pale bands align the operations that execute together.
        for x,w in [(481,297),(1021,297)]:
            c.shape(name+f' overlap band {x}',x,y+133,w,205,'F0F7ED','F0F7ED',width=0)
        if base==0:
            tasks=[(401,80,'Prep.\ni',YELLOW),(481,205,'Read i+1\n→ B',BLUE),
                   (686,92,'Wait',PALE),(778,163,'Post-process\ni',YELLOW),
                   (941,80,'Prep.\ni+1',YELLOW),(1021,205,'Read i+2\n→ A',BLUE),
                   (1226,92,'Wait',PALE),(1318,223,'Post-process\ni+1',YELLOW)]
        else:
            tasks=[(401,80,'Prep.\ni',YELLOW),(481,150,'Write i−1\nfrom B',ROSE),
                   (631,147,'Prefetch\n(optional)',BLUE),(778,163,'Join / validate\ni',YELLOW),
                   (941,80,'Prep.\ni+1',YELLOW),(1021,150,'Write i\nfrom A',ROSE),
                   (1171,147,'Prefetch\n(optional)',BLUE),(1318,223,'Join / validate\ni+1',YELLOW)]
        for j,(x,w,s,col) in enumerate(tasks):
            c.box(name+f' task {j}',x,y+157,w,61,s,col,18)
        for j,x in enumerate([481,1021]):
            s='Decode + parse + op' if base==0 else 'Serialize + compress'
            c.box(name+f' CPE kernel {j}',x,y+272,297,62,
                  s+f'\nbatch {"i / A" if j==0 else "i+1 / B"}',GREEN,20,True)
            c.arrow(name+f' launch {j}',x,y+220,x,y+270)
            c.text(name+f' launch label {j}',x+7,y+234,85,23,'launch',17,align='left')
            c.arrow(name+f' completion {j}',x+297,y+270,x+297,y+220)
            c.text(name+f' completion label {j}',x+302,y+235,139,23,
                   'join / validate' if base==0 else 'complete',17,align='left')
        c.bracket(name+' overlap bracket',481,778,y+352,
                  'Input + CPE work' if base==0 else 'Output / prefetch + CPE work',GREEN_DARK)
        c.text(name+' dependency note',865,y+347,676,47,
               'Post-processing begins only after CPE completion.' if base==0 else
               'Already packed BGZF payloads need compression only.',19,color=GRAY,align='right')
    c.text('Timeline scope',30,866,1535,27,
           'Each panel is a separate schedule. A core group executes one CPE kernel at a time.',19,color=GRAY)
    return c


def exchange_glyph(c,name,x,y,w,h):
    c.shape(name,x,y,w,h,PURPLE,INK,kind='round')
    c.text(name+' title',x+4,y+13,w-8,52,'MPI\nexchange',21,True)
    for j in range(3):
        yy=y+87+j*22
        c.line(name+f' crossed wire {j}',[(x+17,yy),(x+w-17,y+87+(2-j)*22)],VIOLET,1.6,head=True)


def reordering():
    c=Canvas('Distributed sorting and collation',
        'Toy examples on three ranks. Sort numbers denote coordinate keys on one reference; splitters 35/65 '
        'avoid equality ambiguity. Each source supplies a sorted segment to each destination. Output is a '
        'k-way merge, not a fresh full sort of received records. Colors denote destination ranges. '
        'Collate letters are QNAMEs. Example bins: bin(B)=0, bin(A)=2, bin(C)=4, bins=6, ranks=3; '
        'owner=floor(bin*ranks/bins), not bin modulo ranks. Actual order is bin/hash/full QNAME/flag/global order; '
        'same-QNAME groups are contiguous, global QNAME lexical order is not required. Full QNAME resolves hash ties. '
        'Self-owned data remain local. External mode uses algorithmic sorted runs plus bounded exchange/merge; '
        'these temporary runs are independent of optional final rank body spooling. '
        'Sources: swbam-app/algorithms/sort/swbam_sort_optimized_mpi.cpp; '
        'swbam-app/algorithms/collate/swbam_collate_mpi.cpp CollateOwner and CollateMemoryLoserTree.')
    c.section(12,'(a) Coordinate sorting: range partitioning and ordered merge')
    c.shape('Sort background',28,55,1544,337,'F2F7FA',INK,width=1.5)
    headers=[(75,249,'Local records'),(395,246,'Local ordering'),(856,267,'Received segments'),(1251,268,'Merged output')]
    for x,w,s in headers: c.text('Sort '+s,x,72,w,30,s,22,True)
    c.box('Sample splitter plan',396,114,244,34,'Sampled splitters: 35, 65',YELLOW,17)
    raw=[[70,10,60],[30,90,40],[50,20,80]]
    ordered=[[10,60,70],[30,40,90],[20,50,80]]
    recv=[[10,30,20],[60,40,50],[70,90,80]]
    output=[[10,20,30],[40,50,60],[70,80,90]]
    color=lambda v: BLUE if v<35 else YELLOW if v<65 else GREEN
    for r,y in enumerate([171,230,289]):
        c.text(f'Sort rank {r}',42,y,39,38,f'r{r}',19,True)
        c.chips(f'Sort raw {r}',95,y,raw[r],[color(v) for v in raw[r]],cw=63,ch=38,gap=3)
        c.arrow(f'Sort local arrow {r}',300,y+19,389,y+19)
        c.chips(f'Sort ordered {r}',409,y,ordered[r],[color(v) for v in ordered[r]],cw=63,ch=38,gap=3)
        c.arrow(f'Sort into exchange {r}',616,y+19,691,226+(r-1)*32)
        c.arrow(f'Sort out exchange {r}',800,226+(r-1)*32,853,y+19)
        c.text(f'Sort destination {r}',818,y-9,35,19,f'r{r}',16,color=GRAY)
        c.chips(f'Sort segments {r}',870,y,recv[r],color(recv[r][0]),cw=66,ch=38,gap=12)
        c.arrow(f'Sort merge arrow {r}',1116,y+19,1242,y+19)
        c.chips(f'Sort merged {r}',1260,y,output[r],color(output[r][0]),cw=76,ch=38,gap=0)
    exchange_glyph(c,'Sort exchange',692,151,108,170)
    c.text('K way merge label',1122,134,115,24,'k-way merge',18)
    c.text('Sort rank output meaning',59,348,1480,27,
           'Rank-order concatenation: 10, 20, 30 | 40, 50, 60 | 70, 80, 90 → globally coordinate-sorted BAM',20)

    c.section(414,'(b) QNAME collation: bin ownership and contiguous groups')
    c.shape('Collate background',28,457,1544,300,'F1F7EE',INK,width=1.5)
    for x,w,s in [(75,249,'Local records'),(395,246,'Bin / hash ordering'),(856,267,'Owner segments'),(1251,268,'Grouped output')]:
        c.text('Collate '+s,x,472,w,30,s,22,True)
    qcolors={'B':BLUE,'A':YELLOW,'C':GREEN}
    rawq=[['B','A','C'],['C','B','A'],['A','C','B']]
    for r,y in enumerate([525,581,637]):
        c.text(f'Collate rank {r}',42,y,39,38,f'r{r}',19,True)
        c.chips(f'Collate raw {r}',95,y,rawq[r],[qcolors[v] for v in rawq[r]],cw=63,ch=38,gap=3)
        c.arrow(f'Collate sort {r}',300,y+19,389,y+19)
        c.chips(f'Collate ordered {r}',409,y,['B','A','C'],[BLUE,YELLOW,GREEN],cw=63,ch=38,gap=3)
        c.arrow(f'Collate into exchange {r}',616,y+19,691,600+(r-1)*32)
        c.arrow(f'Collate out exchange {r}',800,600+(r-1)*32,853,y+19)
        key=['B','A','C'][r]
        c.chips(f'Collate segments {r}',870,y,[key]*3,qcolors[key],cw=66,ch=38,gap=12)
        c.arrow(f'Collate merge {r}',1116,y+19,1242,y+19)
        c.chips(f'Collate grouped {r}',1260,y,[key]*3,qcolors[key],cw=76,ch=38,gap=0)
    exchange_glyph(c,'Collate exchange',692,515,108,170)
    c.text('Collate merge method',1121,511,116,24,'Loser tree',18)
    c.text('Collate ownership mapping',56,704,1029,27,
           'Example bin owners: B → r0, A → r1, C → r2. Self-owned records stay local.',19)
    c.text('Collate not globally sorted',1126,704,421,27,'Grouped, not globally QNAME-sorted',19,True,color=GREEN_DARK)

    c.text('External mode title',31,786,228,54,'External-memory\nmode',22,True,align='left')
    for i in range(3): c.document(f'Temp sorted run {i}',278+i*37,787+i*5,52,61,str(i+1),PURPLE)
    c.text('Temp runs label',270,858,181,24,'Sorted runs',18)
    for x,w,s in [(477,270,'Bounded exchange'),(799,289,'Streaming multi-way merge'),(1140,285,'Common BAM writer')]:
        c.box('External '+s,x,798,w,58,s,PURPLE if x<1100 else ROSE,20,rounded=True)
    c.arrow('Runs to bounded exchange',411,828,473,828)
    c.arrow('Exchange to streaming merge',747,828,795,828)
    c.arrow('External merge to writer',1088,828,1136,828)
    c.document('External output',1476,790,64,74,'BAM',ROSE)
    c.arrow('External final output',1428,828,1473,828)
    return c


def markdup():
    c=Canvas('One-pass markdup: boundary pre-resolution and ordered retirement',
        'Panel (a) precedes the main stream: read/decode rank head/tail blocks, extract candidate metadata, '
        'allgather starts and collectively verify halo coverage against maximum read length L. '
        'When covered, select boundary candidates, route to coordinate owners, resolve groups and return '
        'source-local decisions H_r. Coverage failure on any rank disables the fast plan; retain boundary '
        'keys conservatively. All edges are schematic; left/right ranks are adjacent coordinate partitions. '
        'The example uses a shared duplicate key k with scores 70 and 60 (illustrative), and not full BAM '
        'record exchange. Paired/single interactions still follow implementation rules. '
        'Panel (b) excludes pre-resolved candidates from ordinary window processing, applies H_r and '
        'local/remote decisions to retained decoded blocks. Optimistic KEEP does not imply immediate output: '
        'active winners and remote decisions may pin records. Only complete front blocks are emitted. '
        'Pending blocks B2/B3 cannot overtake; storage is retired/recycled only after write completion. '
        'One-pass refers to the main record stream; edge probes perform additional limited decoding. '
        'No universal fixed memory bound for arbitrary coverage/fallback is claimed. '
        'Source: swbam-app/algorithms/markdup/swbam_markdup_mpi.cpp '
        'MdBuildOnePassBoundaryPlan, MdOnePassPrepareOutput, one-pass capture/output loop; '
        'swbam_markdup_stream_mpi.cpp candidate-window expiry and boundary retention.')
    c.section(12,'(a) Pre-resolve rank boundaries before the main stream')
    c.panel('Boundary probe region',29,60,608,324,'Probe head / tail BGZF blocks',BLUE)
    c.text('Left probe rank',58,118,245,25,'Rank r−1: tail',22,True)
    c.text('Right probe rank',350,118,247,25,'Rank r: head',22,True)
    c.blocks('Left BGZF blocks',56,163,6,cw=40,ch=43,gap=3,color=PALE,selected=(3,4,5))
    c.blocks('Right BGZF blocks',350,163,6,cw=40,ch=43,gap=3,color=PALE,selected=(0,1,2))
    c.line('Rank partition cut',[(330,144),(330,275)],INK,1.7,dash=True)
    c.arrow('Coordinate progress axis',65,260,603,260,GRAY)
    for x,s in [(227,'b−L'),(330,'b'),(455,'b+L')]:
        c.line('Coordinate tick '+s,[(x,254),(x,267)],GRAY,1.3)
        c.text('Coordinate label '+s,x-37,272,74,25,s,19)
    c.bracket('Probe halo extent',220,465,235,'Boundary halo',GREEN_DARK,False)
    c.text('Probe extraction note',60,317,544,46,
           'CPE decoding + candidate extraction\nExchange candidate metadata, not full BAM records',18)
    c.arrow('Probe to collective coverage',637,218,686,218)
    c.shape('All ranks covered diamond',687,162,166,112,YELLOW,INK,kind='diamond')
    c.text('Coverage check text',707,187,126,63,'Covered\non all ranks?',19,True)
    c.arrow('Covered to owner resolution',853,218,904,218)
    c.text('Coverage yes',856,177,48,25,'Yes',18,color=GREEN_DARK)

    c.panel('Boundary owner resolution',908,92,321,248,'Coordinate owner',GREEN)
    c.box('Left candidate',929,151,277,50,'r−1: key k  |  score 70',BLUE,19)
    c.box('Right candidate',929,215,277,50,'r: key k  |  score 60',ROSE,19)
    c.text('Resolve same duplicate key',930,288,276,29,'Same key → select winner',19,True)
    c.arrow('Return boundary decisions',1229,218,1284,218)
    c.text('Return label',1234,177,53,25,'Return',16)
    c.panel('Boundary decisions',1287,92,284,248,'Cached decisions Hᵣ',YELLOW)
    c.box('Pre resolved keep',1308,151,241,50,'r−1   :   KEEP',GREEN,21,True)
    c.box('Pre resolved duplicate',1308,215,241,50,'r      :   DUP',ROSE,21,True)
    c.text('Boundary decision indexing',1306,288,244,29,'By source record ID',18)
    c.line('Coverage failure fallback',[(770,274),(770,367),(866,367)],GOLD,1.8,head=True,dash=True)
    c.text('Failure label',772,288,68,25,'No',18,color=GOLD)
    c.text('Conservative fallback',877,347,662,48,
           'Otherwise: retain boundary candidates\nin the conservative window',19,color=GOLD,align='left')

    c.section(416,'(b) Retain decoded blocks; emit only the finalized prefix')
    c.panel('Main read and candidate extraction',30,510,269,110,'BAM read (CPE)',BLUE)
    c.text('Decode extraction text',44,554,240,55,'Decode + extract\nDecoded B / candidates C',20)
    c.arrow('Local candidate handoff',299,562,407,562)
    c.text('Local candidate label',344,571,61,22,'Local C',16)
    c.panel('Active candidate window',407,510,376,110,'Candidate window Wᵣ',GREEN)
    c.text('Owner candidate inputs',425,550,338,22,'Local + received candidates',17,color=GREEN_DARK)
    c.box('Paired window state',425,578,158,30,'Paired keys',WHITE,19)
    c.box('Single window state',604,578,158,30,'Single keys',WHITE,19)
    c.text('Exclude already pre resolved',358,448,453,24,'Exclude candidates already resolved by Hᵣ',18,color=GRAY)
    c.panel('Remote owners',949,510,257,110,'Remote owners',PURPLE)
    c.text('Remote owner action',962,554,232,53,'Resolve remote keys\nReturn source decisions',18)
    c.line('Remote candidate routing',[(337,562),(337,479),(1078,479),(1078,508)],VIOLET,1.8,head=True)
    c.text('Remote candidates label',907,448,232,24,'Remote candidates',18,color=VIOLET)
    c.line('Remote decisions apply',[(1078,620),(1078,674)],VIOLET,1.8,head=True,dash=True)
    c.text('Remote decision label',1090,640,212,24,'Returned decisions',17,color=VIOLET,align='left')
    c.box('Apply cached H',1274,510,296,100,'Apply cached Hᵣ\nKEEP / DUP',YELLOW,22,True,rounded=True)

    c.text('Decoded retention label',99,640,226,47,'Retain decoded blocks\n(provisional KEEP)',18)
    c.line('Decoded block storage path',[(80,620),(80,713),(347,713)],BLUE_DARK,2.2,head=True)
    c.text('Queue heading',351,641,371,28,'Pending queue Qᵣ (input order)',21,True,align='left')
    c.text('Queue tail label',1452,641,95,27,'Tail',18,color=GRAY,align='right')
    c.shape('Pending queue outline',347,680,1222,82,WHITE,INK,kind='round',dash=True)
    states=[('B₀','Ready',GREEN),('B₁','Ready',GREEN),('B₂','Pending',YELLOW),('B₃','Ready, held',PALE)]
    for i,(block,state,col) in enumerate(states):
        x=363+i*300
        c.box('Queue block '+block,x,694,272,52,block+'   '+state,col,22,True)
        if i<3:c.arrow('Queue order '+str(i),x+272,720,x+300,720,GRAY)
    c.line('Window decisions apply',[(579,620),(579,630),(765,630),(765,674)],GREEN_DARK,1.8,head=True,dash=True)
    c.text('Local decisions label',777,641,182,23,'Local decisions',17,color=GREEN_DARK,align='left')
    c.line('Cache decisions apply',[(1422,610),(1422,674)],GOLD,1.8,head=True,dash=True)
    c.bracket('Ready output prefix',363,935,773,'Emit the consecutive ready prefix',GREEN_DARK)
    c.text('Queue stop explanation',961,776,593,25,'Stop at the first unresolved block',19,True,color=GOLD)
    c.arrow('Ready blocks to output',470,799,470,823,GREEN_DARK)
    for x,w,s,col in [(347,247,'Patch DUP flags',YELLOW),(633,272,'CPE BGZF compression',GREEN),
                       (944,266,'Append rank body',ROSE),(1249,320,'Retire / recycle storage',BLUE)]:
        c.box('Output '+s,x,825,w,49,s,col,19,rounded=True)
    for x1,x2 in [(594,633),(905,944),(1210,1249)]:c.arrow('Output dependency '+str(x1),x1,850,x2,850)
    c.text('Main stream extra decoding qualifier',34,825,281,49,
           'Boundary probes perform\nlimited extra decoding.',17,color=GRAY)
    return c


def main():
    PREVIEW.mkdir(parents=True,exist_ok=True)
    figures=[overview(),timing(),reordering(),markdup()]
    prs=Presentation();prs.slide_width=Inches(W/120);prs.slide_height=Inches(H/120)
    prs.core_properties.title='SWBAM: Method Figures'
    prs.core_properties.subject='Editable architecture, overlap, reordering and duplicate marking diagrams'
    prs.core_properties.author='SWBAM'
    for i,c in enumerate(figures,1):
        c.validate();c.add_slide(prs)
        svg=c.svg();(PREVIEW/f'{i:02}.svg').write_text(svg,encoding='utf-8')
        cairosvg.svg2png(bytestring=svg.encode(),write_to=str(PREVIEW/f'{i:02}.png'),output_width=2400,output_height=1350)
        print(f'{i}: {c.title}: {len(c.items)} editable scene objects')
    dest=OUT/'swbam_method_figures_redesigned.pptx';prs.save(dest)
    check=Presentation(dest)
    assert len(check.slides)==4
    def flatten(shapes):
        for shape in shapes:
            if shape.shape_type==6:
                yield from flatten(shape.shapes)
            else:
                yield shape
    for c,slide in zip(figures,check.slides):
        shapes=list(flatten(slide.shapes))
        assert all(sh.shape_type!=13 for sh in shapes),'Raster object found'
        expected=[it['text'] for it in c.items if it['type']=='text']
        actual=[sh.text for sh in shapes if sh.has_text_frame and sh.text]
        assert expected==actual,'Editable label mismatch'
    print(dest)


if __name__=='__main__':
    main()
