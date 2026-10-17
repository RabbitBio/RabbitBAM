#!/usr/bin/env python3
"""Redraw the shared six-rank sort/collate diagram as editable PPT shapes.

Uses the scene/export helpers in swbam_method_redesign.py. The existing deck is
left untouched; the third slide is replaced only in a new v2 copy. Dependencies:
python-pptx, Pillow, cairosvg. Run from any directory.
"""

from copy import deepcopy
from pathlib import Path
import sys

sys.dont_write_bytecode = True

import cairosvg
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE
from pptx.util import Inches

from swbam_method_redesign import (
    Canvas, W, H, INK, GRAY, LIGHT, BLUE, BLUE_DARK, GREEN,
    GREEN_DARK, YELLOW, GOLD, PURPLE, WHITE,
)


OUT = Path(__file__).resolve().parent
PREVIEW = Path('/tmp/swbam-redesign/sort-collate-v2')
OWNER_COLORS = ['CBEAF7', 'DAEBCB', 'FFF5C2', 'FAD9BF', 'DCD6F0', 'F4CFD5']


def segment(c, name, x, y, color, local=False):
    """A record segment, not one BAM record or one fixed-size BGZF block."""
    c.shape(name, x, y, 34, 28, color, INK, width=2.4 if local else 1.2)
    for j in range(2):
        c.line(name + f' / records {j}',
               [(x + 6, y + 9 + j * 9), (x + 28, y + 9 + j * 9)],
               INK, 1.1)


def disk_stack(c, name, x, y, label):
    for i in range(2, -1, -1):
        c.shape(name + f' / file {i}', x + i * 6, y - i * 5,
                53, 57, WHITE if i else PURPLE, INK,
                kind='paper', width=1.3)
    c.text(name + ' / description', x - 39, y - 39, 147, 24,
           label, 17, True)
    for j in range(3):
        c.line(name + f' / data {j}',
               [(x + 9, y + 13 + j * 10), (x + 38, y + 13 + j * 10)],
               INK, 1.2)


def draw_shared_path(c):
    c.section(12, '(a) Shared distributed reordering: six MPI ranks')
    c.shape('Shared six-rank processing region', 28, 57, 1544, 430,
            WHITE, INK, kind='round', width=1.8, dash=True)

    # Reference-style stage bands keep the repeated rank rows visually aligned.
    stages = [
        (47, 290, 'Read / extract', 'CPE', GREEN),
        (370, 254, 'Local ordering', 'MPE', YELLOW),
        (650, 180, 'Redistribute', 'by owner', BLUE),
        (853, 298, 'Received segments', 'ordered within each segment', BLUE),
        (1211, 341, 'Merge / BGZF output', 'MPE / CPE', GREEN),
    ]
    for x, w, title, subtitle, color in stages:
        c.shape(title + ' / band', x, 73, w, 57, color, INK, width=1.4)
        c.text(title + ' / title', x + 6, 77, w - 12, 26, title, 22, True)
        c.text(title + ' / subtitle', x + 6, 104, w - 12, 23,
               subtitle, 17)

    for source in range(6):
        c.text(f'Source segment heading {source}', 891 + source * 42,
               136, 34, 22, f's{source}', 16, color=GRAY)

    permutations = [
        [4, 1, 5, 0, 3, 2], [2, 5, 3, 1, 0, 4], [3, 0, 4, 2, 5, 1],
        [1, 3, 0, 5, 2, 4], [5, 2, 1, 4, 0, 3], [0, 4, 2, 3, 1, 5],
    ]
    for r in range(6):
        y = 166 + r * 43
        c.text(f'Source rank {r}', 47, y - 1, 41, 31, f'r{r}', 20, True)
        c.chips(f'Unordered rank {r}', 95, y, permutations[r],
                [OWNER_COLORS[v] for v in permutations[r]],
                cw=34, ch=28, gap=4, size=18)
        c.arrow(f'Local ordering rank {r}', 325, y + 14, 370, y + 14,
                width=2)
        c.chips(f'Owner-ordered rank {r}', 380, y, range(6), OWNER_COLORS,
                cw=34, ch=28, gap=4, size=18)

        c.text(f'Destination rank {r}', 845, y - 1, 41, 31,
               f'r{r}', 20, True)
        for s in range(6):
            segment(c, f'Rank {r} segment from source {s}',
                    891 + s * 42, y, OWNER_COLORS[r], local=(s == r))
        c.arrow(f'Merge rank {r}', 1143, y + 14, 1211, y + 14, width=2)
        for j in range(6):
            c.shape(f'Rank {r} merged record group {j}',
                    1225 + j * 29, y, 29, 28, OWNER_COLORS[r], INK, width=1.2)
        c.arrow(f'Compress rank {r}', 1407, y + 14, 1446, y + 14, width=2)
        c.box(f'Rank body {r}', 1454, y, 85, 28,
              f'body {r}', WHITE, 17, stroke=INK)

    # Brackets denote a many-rank transpose, not a centralized MPI coordinator.
    c.line('All sender ranks bracket', [(627,160), (638,160), (638,414), (627,414)],
           INK, 2)
    c.line('All receiver ranks bracket', [(837,160), (826,160), (826,414), (837,414)],
           INK, 2)
    c.text('Chunked exchange caption', 649, 189, 166, 78,
           'Chunked\npeer-to-peer\nexchange', 21, True)
    c.arrow('Distributed transpose', 638, 283, 826, 283,
            BLUE_DARK, width=4.2)
    c.line('Self-owned local path', [(678,313), (678,331), (793,331), (793,313)],
           INK, 1.8, head=True)
    c.text('Self-owned local caption', 650, 345, 170, 49,
           'Self-owned:\nkeep local', 19)

    c.text('Destination color legend', 49, 443, 568, 25,
           'Cell color / number: destination rank', 19, align='left')
    c.text('Source and local segment legend', 668, 432, 484, 47,
           's0 ... s5: source ranks\nBold segment outline: self-owned data', 17)
    c.text('Output body clarification', 1212, 443, 339, 25,
           'One compressed body per rank', 18)


def draw_partition_policies(c):
    c.panel('Partition policy inset', 28, 512, 616, 373,
            '(b) Different partitioning policies', WHITE)

    c.text('Sort policy title', 49, 569, 100, 28, 'Sort', 23, True,
           color=BLUE_DARK, align='left')
    c.text('Sort policy summary', 187, 569, 434, 28,
           'Samples \u2192 5 splitters \u2192 6 ranges', 20)
    for r in range(6):
        x = 57 + r * 92
        c.box(f'Coordinate owner range {r}', x, 634, 92, 35,
              f'r{r}', OWNER_COLORS[r], 20, True)
        if r:
            c.line(f'Splitter tick {r}', [(x,626), (x,674)], INK, 1.5)
            c.text(f'Splitter label {r}', x - 22, 600, 44, 25,
                   f's{r}', 18)
    c.text('Sort coordinate order', 52, 678, 561, 25,
           'Global coordinate order', 20, True, color=BLUE_DARK)
    c.arrow('Increasing coordinate axis', 84, 712, 581, 712, BLUE_DARK, width=1.6)

    c.text('Collate policy title', 49, 731, 120, 28, 'Collate', 23, True,
           color=GREEN_DARK, align='left')
    c.text('Collate mapping formula', 190, 731, 430, 28,
           'QNAME \u2192 hash / bin \u2192 owner', 20)
    for r in range(6):
        x = 57 + r * 92
        c.text(f'Bin range {r}', x, 766, 92, 24,
               f'{r * 11}\u2013{r * 11 + 10}', 18)
        c.box(f'Collate owner range {r}', x, 794, 92, 32,
              f'r{r}', OWNER_COLORS[r], 20, True)
    c.text('Collate group semantics', 50, 835, 573, 45,
           'Example: 66 bins; same-QNAME groups are contiguous\nNo global QNAME ordering',
           18, color=GREEN_DARK)


def draw_memory_modes(c):
    c.panel('Memory strategy inset', 664, 512, 908, 373,
            '(c) Per-rank memory and external-memory modes', WHITE)

    c.text('Memory mode title', 684, 570, 112, 49, 'Memory\nmode', 21, True,
           color=BLUE_DARK)
    for x, w, s in [(812, 187, 'Ordered data\nin RAM'),
                    (1045, 187, 'Received segments\nin RAM'),
                    (1278, 269, 'Final merge\n+ BGZF compression')]:
        c.box('Memory ' + s, x, 565, w, 66, s,
              BLUE if x < 1250 else GREEN, 19)
    c.arrow('Memory owner exchange', 999, 599, 1045, 599, width=2)
    c.arrow('Memory merge input', 1232, 599, 1278, 599, width=2)

    c.line('Separate storage strategies', [(683,650), (1551,650)], LIGHT, 1.3, dash=True)
    c.text('External mode title', 684, 663, 199, 27, 'External mode', 21, True,
           color=GOLD, align='left')
    c.text('External budget description', 924, 663, 621, 27,
           'Bounded working set; F = merge fan-in', 18, color=GRAY, align='right')

    c.box('Bounded run generation', 694, 735, 138, 60,
          'Fill arena\n+ local sort', YELLOW, 19)
    disk_stack(c, 'Temporary local runs', 867, 737, 'Temp. runs')
    c.box('Bounded external exchange', 957, 735, 142, 60,
          'Chunked\nexchange', BLUE, 19)
    disk_stack(c, 'Temporary received segments', 1134, 737, 'Temp. segments')
    c.box('Bounded multiway merging', 1224, 735, 174, 60,
          'Bounded\nF-way merge', YELLOW, 20)
    c.box('Final external BGZF output', 1433, 735, 113, 60,
          'BGZF\noutput', GREEN, 20)
    for name, x1, x2 in [('Spill local runs', 832, 863),
                          ('Load runs', 932, 953),
                          ('Store received segments', 1099, 1130),
                          ('Read segments', 1199, 1220),
                          ('Final merged output', 1398, 1429)]:
        c.arrow(name, x1, 765, x2, 765, width=2)

    c.line('Repeated bounded run generation',
           [(892,794), (892,816), (763,816), (763,798)],
           GOLD, 1.7, head=True)
    c.text('Repeat run generation caption', 696, 830, 259, 24,
           'Repeat for each input chunk', 17, color=GOLD)

    c.line('Additional intermediate merge passes',
           [(1311,795), (1311,816), (1160,816), (1160,797)],
           GOLD, 1.7, head=True)
    c.text('Additional merge passes caption', 1074, 830, 344, 24,
           'Intermediate passes if segments > F', 17, color=GOLD)
    c.text('Final pass annotation', 1401, 703, 146, 23,
           'Final pass', 17, color=GREEN_DARK)


def build_figure():
    c = Canvas('Sort and collate: shared distributed reordering',
        'Six MPI ranks are shown explicitly. Panel (a) combines the conceptual '
        'organization of sort and collate, not their separate implementations. '
        'Colors / cell numbers denote destination owners; s0-s5 denote source '
        'ranks. Glyphs represent record groups or ordered segments, not a '
        'literal block count. External mode can produce multiple segments per '
        'source. The broad exchange arrow is a distributed data transpose, not '
        'a coordinator or rank0 gather. Self-owned data avoid network transfer; '
        'external mode may still pack/copy/store them. The output is a rank-local '
        'compressed body, not an independent full BAM file per rank. Final '
        'distributed layout supplies the common header and one EOF.\n\n'
        'Sort selects coordinate splitters from samples and produces globally '
        'coordinate-sorted output. Collate uses bin/hash/QNAME ordering and '
        'owner=floor(bin*nranks/nbins): with 66 bins and six ranks, rank r owns '
        'bins 11r through 11r+10. Full QNAME comparisons disambiguate hashes. '
        'The example owner ranges do not imply global lexicographic QNAME order. '
        'Local sorting and merge/exchange control execute on MPE; CPE performs '
        'decode/extract and final compression.\n\n'
        'The external path generates bounded local sorted runs, exchanges '
        'chunks into temporary received segments, and uses memory-budgeted '
        'merge fan-in. Intermediate consolidation may require additional '
        'passes before final merge/compression. The drawing abstracts resident '
        'run optimizations; it does not claim every run must spill. Temporary '
        'algorithm runs/segments differ from an optional final RankBodySink '
        'spool. The -m working-set budget is not a whole-process RSS limit.\n\n'
        'Code anchors: swbam-app/algorithms/collate/swbam_collate_mpi.cpp '
        '(CollateOwner, CollateCompare, CollateExchangeMemory, '
        'CollateExchangeRuns, OptimizedBamExternalCollateMPIImpl); '
        'swbam-app/algorithms/sort/swbam_sort_optimized_mpi.cpp '
        '(OptimizedBamExternalSortMPIImpl, MpiSortChooseSplitters, '
        'MpiSortStrictExchangeRuns, MpiSortStrictConsolidate). '
        'Visual reference: overleaf/reference/plot.pptx, especially its '
        'cyan/green/yellow regions, dark outlines, data arrays and arrows.')
    draw_shared_path(c)
    draw_partition_policies(c)
    draw_memory_modes(c)
    c.validate()
    return c


def all_shapes(shapes):
    for shape in shapes:
        if shape.shape_type == MSO_SHAPE_TYPE.GROUP:
            yield from all_shapes(shape.shapes)
        else:
            yield shape


def validate_editable(slide, expected):
    leaves = list(all_shapes(slide.shapes))
    assert all(s.shape_type != MSO_SHAPE_TYPE.PICTURE for s in leaves)
    actual_text = [s.text for s in leaves if s.has_text_frame and s.text]
    assert actual_text == [it['text'] for it in expected.items if it['type'] == 'text']


def main():
    PREVIEW.mkdir(parents=True, exist_ok=True)
    figure = build_figure()
    single = Presentation()
    single.slide_width = Inches(W / 120)
    single.slide_height = Inches(H / 120)
    single.core_properties.title = figure.title
    single.core_properties.author = 'SWBAM'
    slide = figure.add_slide(single)
    validate_editable(slide, figure)
    single_path = OUT / 'swbam_sort_collate_reordering.pptx'
    single.save(single_path)

    # The figure contains only native drawing elements: no media/hyperlink
    # relationships need to be copied when replacing the third slide.
    deck = Presentation(OUT / 'swbam_method_figures_redesigned.pptx')
    assert len(deck.slides) == 4
    assert deck.slide_width == single.slide_width and deck.slide_height == single.slide_height
    target = deck.slides[2]
    for shape in list(target.shapes):
        target.shapes._spTree.remove(shape._element)
    for shape in slide.shapes:
        target.shapes._spTree.insert_element_before(deepcopy(shape._element), 'p:extLst')
    target.notes_slide.notes_text_frame.text = slide.notes_slide.notes_text_frame.text
    validate_editable(target, figure)
    deck_path = OUT / 'swbam_method_figures_redesigned_v2.pptx'
    deck.save(deck_path)

    svg = figure.svg()
    (PREVIEW / 'sort_collate.svg').write_text(svg, encoding='utf-8')
    cairosvg.svg2png(bytestring=svg.encode('utf-8'),
                    write_to=str(PREVIEW / 'sort_collate.png'),
                    output_width=2400, output_height=1350)
    (OUT / 'swbam_sort_collate_reordering.png').write_bytes(
        (PREVIEW / 'sort_collate.png').read_bytes())
    for path in (single_path, deck_path):
        checked = Presentation(path)
        validate_editable(checked.slides[0 if path == single_path else 2], figure)
        print(f'{path}: {len(checked.slides)} slide(s), native editable elements verified')
    print(f'Preview: {PREVIEW / "sort_collate.png"}')


if __name__ == '__main__':
    main()
