#!/usr/bin/env python3
"""Build the editable SWBAM architecture figure and matching vector/raster exports.

Dependencies: python-pptx, cairosvg, pillow.
Run: python3 swbam_overview.py
The PPTX contains native grouped shapes, editable text and connectors, not a bitmap.
All exports use the same scene; regenerate only before editing the PPTX by hand.
"""

from dataclasses import dataclass
from html import escape
from pathlib import Path
import math
import subprocess

import cairosvg
from PIL import ImageFont
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, MSO_AUTO_SIZE, PP_ALIGN
from pptx.enum.dml import MSO_LINE_DASH_STYLE
from pptx.oxml.xmlchemy import OxmlElement
from pptx.util import Inches, Pt


OUT = Path(__file__).resolve().parent
WIDTH, HEIGHT = 1440, 900
INK = "222B33"
MUTED = "4E5B64"
LINE = "6F7B84"
READ = ("E4EFFA", "2873AC")
WRITE = ("E1F2EA", "217957")
MPE = ("FFF0DC", "AB691B")
NEUTRAL = ("F2F4F6", "737D86")
PURPLE = ("EEE9F6", "735396")
SCENE = []


@dataclass
class Label:
    text: str
    size: float = 22
    bold: bool = False
    color: str = INK


def text(name, x, y, w, h, content, size=22, bold=False,
         color=INK, align="left"):
    SCENE.append(dict(kind="text", name=name, x=x, y=y, w=w, h=h,
                      label=Label(content, size, bold, color), align=align))


def rect(name, x, y, w, h, fill="FFFFFF", stroke=LINE,
         radius=0, dashed=False, stroke_width=1.5):
    SCENE.append(dict(kind="rect", name=name, x=x, y=y, w=w, h=h,
                      fill=fill, stroke=stroke, radius=radius,
                      dashed=dashed, stroke_width=stroke_width))


def node(name, x, y, w, h, title, detail="", palette=NEUTRAL,
         title_size=23, detail_size=20):
    labels = [Label(title, title_size, True)]
    if detail:
        labels.append(Label(detail, detail_size, False, MUTED))
    SCENE.append(dict(kind="node", name=name, x=x, y=y, w=w, h=h,
                      fill=palette[0], stroke=palette[1], radius=7,
                      dashed=False, stroke_width=1.7, labels=labels))


def arrow(name, x1, y1, x2, y2, dashed=False, color=INK, head=True):
    SCENE.append(dict(kind="arrow", name=name, x1=x1, y1=y1,
                      x2=x2, y2=y2, dashed=dashed, color=color, head=head))


def build_scene():
    text("Panel A", 28, 16, 900, 38, "(a) SWBAM architecture", 27, True)
    rect("Application layer", 28, 69, 1384, 83, "F4F3F7", "8C8495")
    text("Application heading", 49, 82, 285, 32, "SWBAM applications", 25, True)
    text("Application interface role", 49, 116, 290, 25, "Batch-level composition", 21)
    text("Streaming applications", 377, 82, 444, 32,
         "Statistics / filtering / conversion", 23, True)
    text("Streaming state", 377, 116, 444, 23,
         "Batch-local processing and aggregation", 21)
    text("Stateful applications", 867, 82, 520, 32,
         "Sort / collate / fixmate / markdup", 23, True)
    text("Stateful responsibilities", 867, 116, 520, 23,
         "Ordering / grouping / cross-record state / MPI", 21)
    arrow("Uses library", 720, 152, 720, 188, dashed=True, color=LINE)
    text("Uses batch interfaces", 741, 159, 350, 25,
         "Batch-level interfaces", 21, color=MUTED)

    rect("SWBAM library boundary", 28, 192, 1384, 347,
         stroke="79838B", dashed=True)
    text("Library heading", 49, 204, 285, 34, "SWBAM library", 25, True)
    text("Library rank scope", 970, 208, 420, 28,
         "Reusable execution within each MPI rank", 21, color=MUTED, align="right")

    node("BAM read pipeline", 49, 254, 306, 84, "BAM read",
         "CPE: decode + parse + op", READ, detail_size=21)
    node("BAM write pipeline", 394, 254, 306, 84, "BAM write",
         "CPE: serialize + compress", WRITE, detail_size=21)
    node("SAM read pipeline", 739, 254, 306, 84, "SAM read",
         "CPE: parse text into records", READ, detail_size=21)
    node("SAM write pipeline", 1084, 254, 306, 84, "SAM write",
         "CPE: format records as text", WRITE, detail_size=21)

    text("Runtime heading", 49, 382, 196, 31, "Batch runtime", 23, True)
    node("MPE preparation", 249, 370, 245, 65, "MPE preparation",
         "Buffers + parameters", MPE, title_size=21, detail_size=21)
    arrow("Launch", 494, 401, 543, 401)
    node("CPE execution", 543, 370, 211, 65, "CPE execution",
         "Selected kernel", PURPLE, title_size=21, detail_size=21)
    arrow("Complete and join", 754, 401, 803, 401)
    node("MPE post-processing", 803, 370, 305, 65, "MPE post-processing",
         "Validate batch, then process", MPE, title_size=22, detail_size=21)
    text("Buffer reuse", 1137, 372, 251, 29,
         "Reusable buffers", 21, True, align="center")
    text("Pipeline overlap", 1137, 404, 251, 25,
         "Overlap I/O + CPE work", 21, align="center")

    text("I/O heading", 49, 472, 235, 32, "I/O and MPI support", 22, True)
    node("Input and partitioning", 306, 459, 517, 63,
         "Input backend + rank partitioning", "Memory / POSIX / MPI-IO",
         NEUTRAL, title_size=21, detail_size=21)
    node("Output and layout", 863, 459, 527, 63,
         "Rank body sink + distributed output",
         "Memory / spool  \u2192  Memory / MPI-IO",
         NEUTRAL, title_size=21, detail_size=21)

    text("Panel B", 28, 559, 1200, 38,
         "(b) Example batch-level compositions", 27, True)
    # All three stages are sequential dependencies, not concurrent CPE kernels.
    examples = [
        ("Record counting", 49, 610, [
            ("BAM read", "Count (CPE)", READ),
            ("Aggregate", "MPE", MPE),
            ("MPI reduce", "Global count", NEUTRAL)]),
        ("BAM filtering / rewriting", 753, 610, [
            ("BAM read", "Filter / retain", READ),
            ("Packing plan", "MPE", MPE),
            ("BAM write", "Encode (CPE)", WRITE)]),
        ("BAM-to-SAM conversion", 49, 749, [
            ("BAM read", "Decode + parse", READ),
            ("Batch handoff", "MPE", MPE),
            ("SAM write", "Format (CPE)", WRITE)]),
        ("SAM-to-BAM conversion", 753, 749, [
            ("SAM read", "Parse (CPE)", READ),
            ("Packing plan", "MPE", MPE),
            ("BAM write", "Encode (CPE)", WRITE)]),
    ]
    for title, x, y, stages in examples:
        text(title + " heading", x, y, 650, 30, title, 23, True)
        for i, (label, detail, palette) in enumerate(stages):
            xx = x + i * 226
            node(title + " / " + label, xx, y + 39, 188, 72,
                 label, detail, palette, title_size=22, detail_size=21)
            if i < 2:
                arrow(title + " / batch " + str(i),
                      xx + 188, y + 75, xx + 226, y + 75)


def font_path(bold):
    query = "Liberation Sans:style=" + ("Bold" if bold else "Regular")
    return subprocess.check_output(
        ["fc-match", "-f", "%{file}", query], text=True).strip()


def validate_scene():
    fonts = {bold: font_path(bold) for bold in (False, True)}
    for item in SCENE:
        if item["kind"] == "arrow":
            assert 0 <= item["x1"] <= WIDTH and 0 <= item["x2"] <= WIDTH, item["name"]
            assert 0 <= item["y1"] <= HEIGHT and 0 <= item["y2"] <= HEIGHT, item["name"]
            assert (item["x1"], item["y1"]) != (item["x2"], item["y2"]), item["name"]
            continue
        assert item["x"] >= 0 and item["y"] >= 0, item["name"]
        assert item["x"] + item["w"] <= WIDTH, item["name"]
        assert item["y"] + item["h"] <= HEIGHT, item["name"]
        if item["kind"] not in ("text", "node"):
            continue
        labels = item.get("labels", [item.get("label")])
        pad = 20 if item["kind"] == "node" else 0
        for label in labels:
            font = ImageFont.truetype(fonts[label.bold], round(label.size * 4))
            width = font.getlength(label.text) / 4
            assert width <= item["w"] - pad, (
                f"Text overflow: {item['name']}: {label.text}: "
                f"{width:.1f} > {item['w'] - pad}")


def text_parts(item):
    if item["kind"] == "text":
        yield item["label"], item["x"], item["y"], item["w"], item["h"], item["align"]
    else:
        labels = item["labels"]
        heights = [label.size * 1.22 for label in labels]
        gap = 4
        top = item["y"] + (item["h"] - sum(heights) - gap * (len(labels) - 1)) / 2
        for label, h in zip(labels, heights):
            yield label, item["x"] + 8, top, item["w"] - 16, h, "center"
            top += h + gap


def to_svg(title="SWBAM architecture and pipeline composition", description=None):
    # The PDF has a journal-ready width of 178 mm. PNG export overrides this.
    if description is None:
        description = ("Library/application separation and batch-level composition. "
                       "Blue: reading; green: writing; amber: MPE work.")
    lines = [f'<svg xmlns="http://www.w3.org/2000/svg" width="178mm" '
             f'height="{178 * HEIGHT / WIDTH:.5f}mm" viewBox="0 0 {WIDTH} {HEIGHT}">',
             f'<title>{escape(title)}</title>',
             f'<desc>{escape(description)}</desc>',
             f'<rect width="{WIDTH}" height="{HEIGHT}" fill="white"/>']
    for item in SCENE:
        lines.append(f'<g id="{escape(item["name"], quote=True)}">')
        if item["kind"] == "arrow":
            x1, y1, x2, y2 = (item[k] for k in ("x1", "y1", "x2", "y2"))
            ux, uy = x2 - x1, y2 - y1
            norm = math.hypot(ux, uy)
            ux, uy = ux / norm, uy / norm
            head = item.get("head", True)
            bx, by = (x2 - 9 * ux, y2 - 9 * uy) if head else (x2, y2)
            dash = ' stroke-dasharray="5 4"' if item["dashed"] else ""
            lines.append(f'<line x1="{x1}" y1="{y1}" x2="{bx}" y2="{by}" '
                         f'stroke="#{item["color"]}" stroke-width="1.8"{dash}/>')
            if head:
                lines.append(f'<polygon points="{x2},{y2} {bx - 3.7 * uy},{by + 3.7 * ux} '
                             f'{bx + 3.7 * uy},{by - 3.7 * ux}" fill="#{item["color"]}"/>')
        elif item["kind"] in ("rect", "node"):
            dash = ' stroke-dasharray="7 5"' if item["dashed"] else ""
            lines.append(f'<rect x="{item["x"]}" y="{item["y"]}" width="{item["w"]}" '
                         f'height="{item["h"]}" rx="{item["radius"]}" fill="#{item["fill"]}" '
                         f'stroke="#{item["stroke"]}" stroke-width="{item["stroke_width"]}"{dash}/>')
        if item["kind"] in ("text", "node"):
            for label, x, y, w, h, align in text_parts(item):
                anchor = {"left": "start", "center": "middle", "right": "end"}[align]
                xx = x + ({"left": 0, "center": w / 2, "right": w}[align])
                yy = y + h / 2 + label.size * 0.35
                lines.append(f'<text x="{xx:.3f}" y="{yy:.3f}" text-anchor="{anchor}" '
                             f'font-family="Arial, Liberation Sans, sans-serif" '
                             f'font-size="{label.size}" font-weight="{700 if label.bold else 400}" '
                             f'fill="#{label.color}">{escape(label.text)}</text>')
        lines.append('</g>')
    return "\n".join(lines + ['</svg>']) + "\n"


def native_text(shapes, name, label, x, y, w, h, align):
    box = shapes.add_textbox(Inches(x / 100), Inches(y / 100),
                            Inches(w / 100), Inches(h / 100))
    box.name = name
    tf = box.text_frame
    tf.clear()
    tf.word_wrap = False
    tf.auto_size = MSO_AUTO_SIZE.NONE
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    para = tf.paragraphs[0]
    para.alignment = {"left": PP_ALIGN.LEFT, "center": PP_ALIGN.CENTER,
                      "right": PP_ALIGN.RIGHT}[align]
    para.space_before = para.space_after = Pt(0)
    para.line_spacing = 1.0
    run = para.add_run()
    run.text = label.text
    run.font.name = "Arial"
    run.font.size = Pt(label.size * 0.72)
    run.font.bold = label.bold
    run.font.color.rgb = RGBColor.from_string(label.color)
    return box


def to_pptx(path, title="SWBAM architecture and pipeline composition", notes=None):
    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(WIDTH / 100), Inches(HEIGHT / 100)
    prs.core_properties.title = title
    prs.core_properties.subject = "Editable Method figure; all elements are native shapes"
    prs.core_properties.author = "SWBAM"
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = RGBColor(255, 255, 255)
    for item in SCENE:
        if item["kind"] == "arrow":
            shape = slide.shapes.add_connector(
                MSO_CONNECTOR.STRAIGHT,
                Inches(item["x1"] / 100), Inches(item["y1"] / 100),
                Inches(item["x2"] / 100), Inches(item["y2"] / 100))
            shape.name = item["name"]
            shape.line.color.rgb = RGBColor.from_string(item["color"])
            shape.line.width = Pt(1.8 * 0.72)
            if item["dashed"]:
                shape.line.dash_style = MSO_LINE_DASH_STYLE.DASH
            if item.get("head", True):
                end = OxmlElement("a:tailEnd")
                end.set("type", "triangle")
                end.set("w", "sm")
                end.set("len", "med")
                shape._element.spPr.get_or_add_ln().append(end)
            continue
        shapes = slide.shapes
        if item["kind"] == "node":
            group = slide.shapes.add_group_shape()
            group.name = item["name"]
            shapes = group.shapes
        if item["kind"] in ("rect", "node"):
            kind = MSO_SHAPE.ROUNDED_RECTANGLE if item["radius"] else MSO_SHAPE.RECTANGLE
            shape = shapes.add_shape(kind, Inches(item["x"] / 100),
                                     Inches(item["y"] / 100), Inches(item["w"] / 100),
                                     Inches(item["h"] / 100))
            shape.name = item["name"] + " / frame"
            if item["radius"]:
                shape.adjustments[0] = item["radius"] / min(item["w"], item["h"])
            shape.fill.solid()
            shape.fill.fore_color.rgb = RGBColor.from_string(item["fill"])
            shape.line.color.rgb = RGBColor.from_string(item["stroke"])
            shape.line.width = Pt(item["stroke_width"] * 0.72)
            if item["dashed"]:
                shape.line.dash_style = MSO_LINE_DASH_STYLE.DASH
        if item["kind"] in ("text", "node"):
            for i, part in enumerate(text_parts(item)):
                native_text(shapes, item["name"] + f" / label {i + 1}", *part)
    slide.notes_slide.notes_text_frame.text = NOTES if notes is None else notes
    prs.save(path)
    # Ensure the exported presentation remains editable and contains no pictures.
    check = Presentation(path)
    assert len(check.slides) == 1
    def walk(shapes):
        for shape in shapes:
            yield shape
            if hasattr(shape, "shapes"):
                yield from walk(shape.shapes)
    native = list(walk(check.slides[0].shapes))
    assert not any(shape.shape_type == 13 for shape in native)
    expected = [part[0].text for item in SCENE if item["kind"] in ("node", "text")
                for part in text_parts(item)]
    actual = [shape.text for shape in native if shape.has_text_frame and shape.text]
    assert actual == expected
    print(f"PPTX checked: {len(native)} native objects, {len(actual)} editable labels; no bitmap")


NOTES = """Figure: SWBAM architecture and pipeline composition

EDITING
One editable slide. Pipeline and post-processing modules are grouped native
shapes; text and arrows remain editable. Use PowerPoint's Selection Pane to
select named modules. Export the slide as PDF after editing; the PNG/PDF/SVG
provided alongside this file reflect the original generated layout only.

SEMANTICS
Panel (a) is an architecture/capability view, not four consecutive stages.
Panel (b) shows batch-level dependencies, not concurrently running CPE kernels.
Read pipelines are blue, write pipelines green, MPE work amber. Applications
manage algorithm state and algorithm-specific MPI exchanges; the library
provides batch execution, data access, rank partitioning and output layout.
Input is partitioned into BGZF spans (BAM) or record-aligned text ranges (SAM).
BAM read kernels may return records, counters or algorithm-specific metadata.
The CPE work shown is representative of optimized application paths; it does
not assert that every SDK bam1_t materialization runs on CPEs. Decode-only and
compression-only variants are omitted. The implementation assumes BAM records
do not span BGZF blocks. Data retained across batches require owned storage.
The sink storage policy is independent from final output (memory / MPI-IO).
No root-wide full-payload gather is implied for MPI-IO output.
The figure does not claim arbitrary runtime composition of CPE record operators.

CODE REFERENCES (relative to the RabbitBAM project)
swbam-sdk/include/swbam/cpe_pipeline.h: CpeBatchOperator / RunCpeReadPipeline
swbam-sdk/lib/cpe/swbam_cpe_bam_read_pipeline.cpp: prepare, launch, validate, post-process
swbam-sdk/lib/cpe/swbam_cpe_bam_write_pipeline.cpp: current/pending output batches
swbam-sdk/lib/operators/swbam_record_count_operator.cpp: read + accumulation
swbam-sdk/lib/operators/swbam_bam_transform_operator.cpp: filter + packing + writer
swbam-sdk/lib/operators/swbam_bam_to_sam_operator.cpp: read + SAM writer
swbam-sdk/lib/operators/swbam_sam_to_bam_operator.cpp: SAM read + packing + BAM writer
swbam-sdk/include/swbam/mpi_runtime.h: input planning and distributed output

MANUSCRIPT AND VISUAL REFERENCES
overleaf/SWBAM/bare_jrnl_new_sample4.tex: Methods / Overview and Library Design
RabbitBAM, Fig. 1 and Fig. 3: explicit reusable read/write framework blocks
SWBWA, Fig. 1: explicit MPE/CPE roles; style reference only, not its algorithm
"""


CAPTION = r"""% Insert in the Overview subsection. Paths are relative to overleaf/SWBAM/.
\begin{figure*}[!t]
    \centering
    \includegraphics[width=\textwidth]{../figs/swbam_overview.pdf}
    \caption{SWBAM architecture and pipeline composition.
    (a) Applications reuse four batch-based processing pipelines, a common
    MPE--CPE execution runtime, and I/O and MPI support. Application-specific
    grouping, ordering, communication, and cross-record state remain in the
    application layer. (b) Representative operations connect read pipelines,
    MPE-side batch post-processing, and write pipelines or MPI reduction.
    Arrows show batch-level dependencies, not concurrent CPE kernel execution.}
    \label{fig:swbam_overview}
\end{figure*}
"""


def main():
    build_scene()
    validate_scene()
    svg = to_svg()
    (OUT / "swbam_overview.svg").write_text(svg, encoding="utf-8")
    cairosvg.svg2png(bytestring=svg.encode(), output_width=4320,
                    write_to=str(OUT / "swbam_overview.png"))
    cairosvg.svg2pdf(bytestring=svg.encode(), write_to=str(OUT / "swbam_overview.pdf"))
    to_pptx(OUT / "swbam_overview.pptx")
    (OUT / "swbam_overview.tex").write_text(CAPTION, encoding="utf-8")
    print(f"Wrote PPTX, PNG, PDF, SVG and LaTeX snippet to {OUT}")


if __name__ == "__main__":
    main()
