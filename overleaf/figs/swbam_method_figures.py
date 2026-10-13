#!/usr/bin/env python3
"""Generate editable SWBAM method diagrams with matching PNG/PDF/SVG exports.

Run: python3 swbam_method_figures.py
Uses the native PowerPoint/SVG drawing helpers in swbam_overview.py.
Dependencies: python-pptx, cairosvg, pillow; fontconfig for layout checks.
"""

from copy import deepcopy
import sys

import cairosvg
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.util import Inches

sys.dont_write_bytecode = True
import swbam_overview as d


OUT = d.OUT
AMBER = d.MPE
READY = d.WRITE
GRAY = d.NEUTRAL


def title(name, y, content):
    d.text(name, 28, y, 1384, 38, content, 26, True)


def label(name, x, y, w, content, size=21, align="center", color=d.MUTED):
    d.text(name, x, y, w, 28, content, size, color=color, align=align)


def line(name, x1, y1, x2, y2, dashed=False, color=d.LINE):
    d.arrow(name, x1, y1, x2, y2, dashed=dashed, color=color, head=False)


def route(name, points, dashed=False, color=d.INK, head=True):
    for i, (p, q) in enumerate(zip(points, points[1:])):
        d.arrow(f"{name} / {i}", *p, *q, dashed=dashed, color=color,
                head=head and i == len(points) - 2)


def bracket(name, x1, x2, y, text, color=d.LINE, below=True):
    dy = -7 if below else 7
    route(name, [(x1, y + dy), (x1, y), (x2, y), (x2, y + dy)],
          color=color, head=False)
    label(name + " label", x1, y + 6 if below else y - 34,
          x2 - x1, text, color=color)


def lane_task(name, x, y, w, text, sub="", palette=AMBER):
    d.node(name, x, y, w, 65, text, sub, palette,
           title_size=22, detail_size=21)


def overlap():
    title("Read title", 17, "(a) BAM read: next-batch input overlaps current CPE work")
    d.arrow("Read time", 180, 87, 1400, 87, color=d.LINE)
    label("Read time label", 1190, 54, 208, "Time (schematic)")
    d.text("Read MPE lane", 32, 160, 124, 33, "MPE", 26, True)
    d.text("Read CPE lane", 32, 291, 124, 33, "CPEs", 26, True)
    lane_task("Read prepare i", 180, 145, 160, "Prepare i", "input A")
    lane_task("Read i+1", 340, 145, 305, "Read batch i+1", "next input: B", d.READ)
    lane_task("Read join i", 645, 145, 125, "Wait", "join", GRAY)
    lane_task("Read post i", 770, 145, 240, "Post-process i", "after validation")
    lane_task("Read prepare i+1", 1010, 145, 140, "Prepare", "i+1")
    lane_task("Read i+2", 1150, 145, 250, "Read batch i+2", "next input: A", d.READ)
    lane_task("Read kernel i", 340, 275, 430, "Decode + parse + op", "batch i / input A", d.PURPLE)
    lane_task("Read kernel i+1", 1150, 275, 250, "Same kernel", "batch i+1 / input B", d.PURPLE)
    d.arrow("Read launch i", 340, 212, 340, 273)
    label("Read launch label", 345, 230, 103, "launch", 20)
    d.arrow("Read complete i", 770, 273, 770, 212)
    label("Read complete label", 775, 230, 170, "CPE complete", 20)
    d.arrow("Read launch i+1", 1150, 212, 1150, 273)
    bracket("Read overlap", 340, 645, 352, "Input / CPE overlap", d.READ[1])
    label("Read post dependency", 695, 355, 710,
          "Post-processing starts after join and validation.", 22)

    title("Write title", 437, "(b) BAM write: previous output overlaps current compression")
    d.arrow("Write time", 180, 507, 1400, 507, color=d.LINE)
    label("Write time label", 1190, 474, 208, "Time (schematic)")
    d.text("Write MPE lane", 32, 583, 124, 33, "MPE", 26, True)
    d.text("Write CPE lane", 32, 709, 124, 33, "CPEs", 26, True)
    lane_task("Write prepare i", 180, 567, 160, "Prepare i", "output A")
    lane_task("Write previous", 340, 567, 260, "Write batch i-1", "pending output: B", d.WRITE)
    lane_task("Write prefetch", 600, 567, 225, "Prefetch i+1", "optional input", d.READ)
    lane_task("Write validate", 825, 567, 185, "Join / validate", "pending A")
    lane_task("Write prepare next", 1010, 567, 140, "Prepare", "i+1")
    lane_task("Write current", 1150, 567, 250, "Write batch i", "pending output: A", d.WRITE)
    lane_task("Write kernel i", 340, 694, 485, "Serialize + compress", "batch i / output A", d.PURPLE)
    lane_task("Write kernel i+1", 1150, 694, 250, "Same kernel", "batch i+1 / output B", d.PURPLE)
    d.arrow("Write launch i", 340, 634, 340, 692)
    label("Write launch label", 345, 649, 103, "launch", 20)
    d.arrow("Write complete i", 825, 692, 825, 634)
    label("Write complete label", 830, 649, 170, "CPE complete", 20)
    d.arrow("Write launch next", 1150, 634, 1150, 692)
    bracket("Write overlap", 340, 825, 772, "Output / prefetch overlaps CPE work", d.WRITE[1])
    label("Write raw payload option", 870, 776, 530,
          "Raw BGZF payloads: compression only.", 21)
    label("Schedule scope", 40, 853, 1360,
          "Independent read/write schedules for one core group; CPE kernels are not concurrent.", 22)


def markdup():
    title("Boundary panel", 17, "(a) Boundary pre-resolution before the main stream")
    label("Left rank", 70, 61, 640, "Rank r-1", 24, color=d.INK)
    label("Right rank", 710, 61, 640, "Rank r", 24, color=d.INK)
    # Schematic coordinates on one reference: probe extent exceeds b +/- L.
    for i in range(10):
        selected = i in (3, 4, 5, 6)
        palette = d.READ if selected else GRAY
        d.rect(f"Edge BGZF {i}", 70 + i * 128, 110, 124, 52,
               palette[0], palette[1])
        label(f"Edge BGZF label {i}", 70 + i * 128, 123, 124,
              "BGZF", 21, color=d.INK)
    bracket("Tail probe", 454, 706, 174, "Tail probe", d.READ[1])
    bracket("Head probe", 714, 962, 174, "Head probe", d.READ[1])
    line("Rank cut", 710, 100, 710, 239, dashed=True, color=d.INK)
    d.arrow("Coordinate axis", 70, 239, 1352, 239, color=d.LINE)
    for x, text in ((542, "b-L"), (710, "b"), (878, "b+L")):
        line("Tick " + text, x, 229, x, 245, color=d.INK)
        label("Coordinate " + text, x - 55, 247, 110, text, 22, color=d.INK)
    label("Coordinate meaning", 1080, 205, 270, "Genomic coordinate", 21)

    d.node("Probe extract", 40, 301, 248, 70, "Probe + extract",
           "CPE: head/tail blocks", d.READ, title_size=23, detail_size=21)
    d.node("Coverage check", 337, 301, 290, 70, "Coverage check",
           "Collective over all ranks", AMBER, title_size=23, detail_size=21)
    d.node("Resolve boundary", 692, 301, 339, 70, "Select / route / resolve",
           "At coordinate owners", d.PURPLE, title_size=23, detail_size=21)
    d.node("Build boundary plan", 1090, 301, 310, 70, "Boundary plan H_r",
           "Source-local decisions", AMBER, title_size=23, detail_size=21)
    d.arrow("Probe to coverage", 288, 336, 337, 336)
    d.arrow("Covered branch", 627, 336, 692, 336)
    label("Covered branch label", 625, 271, 118, "Covered", 20)
    d.arrow("Return boundary decisions", 1031, 336, 1090, 336)
    label("Return decisions label", 1017, 271, 88, "Return", 20)
    route("Fallback branch", [(483, 371), (483, 400), (532, 400)],
          dashed=True, color=AMBER[1])
    label("Fallback explanation", 540, 386, 850,
          "Not covered: retain boundary candidates in the conservative window.", 21, "left", AMBER[1])
    label("Candidate-only exchange", 33, 391, 420,
          "Exchange metadata, not BAM records.", 21)

    title("Pending panel", 440, "(b) Resolve decisions, then emit only the ready prefix")
    d.node("Cached plan reference", 60, 500, 310, 75, "Cached H_r",
           "Pre-resolved KEEP / DUP", AMBER, title_size=24, detail_size=21)
    d.node("Candidate window", 500, 500, 400, 75, "Candidate window W_r",
           "Local + remote owner decisions", d.PURPLE, title_size=24, detail_size=21)
    d.node("Main batch extraction", 1055, 500, 345, 75, "Decode + extract",
           "CPE: current batch", d.READ, title_size=24, detail_size=21)
    d.arrow("Main candidates", 1055, 537, 900, 537)
    label("Main candidate label", 900, 473, 155, "Candidates", 20)
    label("Drop already resolved", 900, 551, 155,
          "exclude H_r", 20)
    # Decision edges update states; the decoded bytes remain in Q_r.
    d.arrow("Cached decisions to queue", 215, 575, 215, 647,
            dashed=True, color=AMBER[1])
    route("Window decisions to queue", [(700, 575), (700, 622), (885, 622), (885, 647)],
          dashed=True, color=d.PURPLE[1])
    d.arrow("Decoded blocks to tail", 1228, 575, 1228, 647, color=d.READ[1])
    label("Append decoded label", 1240, 600, 162, "Decoded B", 21)
    label("Queue label", 20, 614, 160, "Q_r: head", 21, "left", d.INK)
    blocks = [(60, "B_0", "Ready", READY), (400, "B_1", "Ready", READY),
              (740, "B_2", "Pending", AMBER), (1080, "B_3", "Ready, but blocked", GRAY)]
    for x, block, state, palette in blocks:
        d.node("Queue " + block, x, 651, 290, 75,
               block, state, palette, title_size=24, detail_size=22)
    for x in (350, 690, 1030):
        d.arrow(f"Queue order {x}", x, 688, x + 50, 688, color=d.LINE)
    bracket("Ready prefix", 60, 690, 742, "Consecutive ready prefix", d.WRITE[1])
    label("Blocked suffix", 747, 748, 623,
          "No block may pass the first unresolved block.", 22, color=AMBER[1])

    d.arrow("Prefix to patch", 205, 744, 205, 788, color=d.WRITE[1])
    d.node("Patch duplicate flags", 60, 790, 290, 74,
           "Patch duplicate flags", "MPE: decoded records", AMBER,
           title_size=23, detail_size=21)
    d.arrow("Patch to writer", 350, 827, 400, 827)
    d.node("Compress ready prefix", 400, 790, 290, 74,
           "BAM write pipeline", "CPE: BGZF compression", d.WRITE,
           title_size=23, detail_size=21)
    d.arrow("Writer to sink", 690, 827, 740, 827)
    d.node("Commit ready output", 740, 790, 290, 74,
           "Append rank body", "After CPE completion", d.WRITE,
           title_size=23, detail_size=21)
    d.arrow("Retire block storage", 1030, 827, 1080, 827)
    d.node("Recycle storage", 1080, 790, 290, 74,
           "Retire / recycle", "Decoded block storage", GRAY,
           title_size=23, detail_size=21)


def pipeline_row(prefix, y, stages):
    for i, (heading, detail, palette) in enumerate(stages):
        x = 35 + i * 284
        d.node(prefix + " / " + heading, x, y, 230, 78, heading, detail,
               palette, title_size=22, detail_size=21)
        if i < len(stages) - 1:
            d.arrow(prefix + f" / data {i}", x + 230, y + 39, x + 284, y + 39)


def chips(name, x, y, values, palette, cell_width=92):
    for i, value in enumerate(values):
        d.rect(name + f" / {i}", x + i * cell_width, y, cell_width - 4, 45,
               fill=palette[0], stroke=palette[1])
        label(name + f" / value {i}", x + i * cell_width, y + 8,
              cell_width - 4, value, 23, color=d.INK)


def reordering():
    title("Sort panel", 18, "(a) Coordinate sort: partition by sampled coordinate splitters")
    pipeline_row("Sort", 87, [
        ("BAM read / extract", "Keys + raw records", d.READ),
        ("Local ordering", "Sort metadata", AMBER),
        ("Partition / exchange", "Sample-based ranges", d.PURPLE),
        ("k-way merge", "Ordered segments", AMBER),
        ("BAM write", "BGZF compression", d.WRITE),
    ])
    label("Sort toy caption", 39, 187, 1360,
          "Illustrative result after merging (one chromosome; coordinate splitter at 25)", 22)
    label("Sort rank0", 70, 240, 160, "Rank 0", 24, "left", d.INK)
    chips("Sort rank0 records", 248, 233, ["10", "20", "25"], d.READ, 95)
    label("Sort range order", 564, 239, 156, "max < min", 21, color=d.INK)
    label("Sort rank1", 744, 240, 155, "Rank 1", 24, "left", d.INK)
    chips("Sort rank1 records", 913, 233, ["70", "75", "90"], d.WRITE, 95)
    label("Sort range0", 195, 285, 395, "Lower coordinate range", 21)
    label("Sort range1", 861, 285, 401, "Higher coordinate range", 21)
    label("Sort property", 35, 323, 1360,
          "Concatenating rank bodies preserves global coordinate order.", 22, color=d.INK)

    title("Collate panel", 386, "(b) Collate: partition by QNAME hash bins")
    pipeline_row("Collate", 455, [
        ("BAM read / extract", "Hash + raw records", d.READ),
        ("Local ordering", "Bin / hash / QNAME", AMBER),
        ("Partition / exchange", "Bin owner; chunks", d.PURPLE),
        ("Loser-tree merge", "Contiguous groups", AMBER),
        ("BAM write", "BGZF compression", d.WRITE),
    ])
    label("Collate toy caption", 40, 551, 1350,
          "Illustrative bin ownership: qB belongs to rank 0; qA and qC belong to rank 1", 22)
    label("Collate rank0", 70, 605, 160, "Rank 0", 24, "left", d.INK)
    chips("Collate rank0 records", 248, 598, ["qB", "qB"], d.READ, 95)
    label("Collate rank1", 744, 605, 155, "Rank 1", 24, "left", d.INK)
    chips("Collate rank1 records", 913, 598, ["qA", "qA", "qC", "qC"], d.WRITE, 95)
    bracket("Qname B group", 248, 434, 657, "One QNAME group", d.READ[1])
    bracket("Qname A group", 913, 1099, 657, "qA group", d.WRITE[1])
    bracket("Qname C group", 1103, 1289, 657, "qC group", d.WRITE[1])
    label("Collate property", 33, 698, 1360,
          "Equal QNAMEs stay together; global lexicographic QNAME order is not required.", 22, color=d.INK)

    title("External panel", 745, "(c) External mode: preserve the partition rule, bound the working set")
    xs = [35, 390, 745, 1100]
    stages = [("Generate ordered runs", "Temporary run storage", GRAY),
              ("Bounded exchange", "Owner-directed chunks", d.PURPLE),
              ("Streaming merge", "Bounded merge fan-in", AMBER),
              ("Same BAM writer", "Rank-local output body", d.WRITE)]
    for i, (heading, detail, palette) in enumerate(stages):
        d.node("External / " + heading, xs[i], 799, 302, 78,
               heading, detail, palette, title_size=23, detail_size=21)
        if i < 3:
            d.arrow("External transfer " + str(i), xs[i] + 302, 838, xs[i + 1], 838)


COMMON_NOTES = """EDITING AND SCOPE
All text, nodes, block strips and arrows are editable PowerPoint objects.
Module frames and labels are grouped. Use the Selection Pane for named objects.
PNG/PDF/SVG are matching exports from the original scene; after manual PPTX
edits, export a new PDF/PNG in PowerPoint. The code and LaTeX manuscript are
not modified by this figure generator. These diagrams show current supported
BAM workflows under the block-aligned record assumption.
"""

FIGURES = [
    dict(stem="swbam_batch_overlap", build=overlap,
         title="Batch-level overlap between MPE and CPE execution",
         caption=r"""Batch-level MPE--CPE overlap within one core group.
    (a) With read-side input overlap enabled, the MPE reads the next batch into
    the alternate input buffer while the CPEs process the current batch;
    batch post-processing follows CPE completion and validation.
    (b) With write-side output overlap enabled, the MPE emits a previously
    compressed batch while the CPEs process the current batch. Optional
    next-input prefetch is shown. Bars are schematic, not measured durations;
    the two panels are separate schedules, not concurrent read/write kernels.""",
         notes="""CODE BASIS
swbam-sdk/lib/cpe/swbam_cpe_bam_read_pipeline.cpp: RunCpeReadPipeline
  Prepare -> spawn -> ReadNext(next) -> DuringKernel -> join -> Validate
  -> PostProcessBatch -> swap. Panel (a) depicts overlap_input=true and
  prefetch_during_post_process=false. Other operator options exist.
swbam-sdk/lib/cpe/swbam_cpe_bam_write_pipeline.cpp:
  Start prepares output, spawns CPE work, then FlushPending() emits the
  preceding batch. SubmitWithPrefetch may prepare the next source batch
  before Complete() joins and validates. Completed output becomes pending.
  Panel (b) depicts steady state with overlap_output=true; optional source
  prefetch is shown. The first batch has no previous output to flush.
  Input A/B in (a) and output A/B in (b) are separate buffer pairs.
  Record input uses serialization + compression; raw payload input only
  uses compression. Write denotes appending to the configured output sink,
  not necessarily writing to physical disk. The final pending batch is
  flushed at termination (not drawn). Batch durations are illustrative.
"""),
    dict(stem="swbam_markdup_boundary", build=markdup,
         title="Boundary resolution and ordered emission in distributed markdup",
         caption=r"""Boundary resolution and ordered emission in one-pass
    distributed duplicate marking. (a) Head/tail BGZF blocks are probed before
    the main stream. For a boundary $b$ on the same reference, the probe must
    cover the required $b-L$ and $b+L$ limits; coverage is checked collectively.
    Covered boundary candidates are routed to their coordinate owners and
    resolved, with source-local decisions cached in $H_r$. Insufficient
    coverage selects conservative boundary retention. (b) Pre-resolved
    candidates are excluded from ordinary window processing. Decisions from
    $H_r$ and the candidate window update retained decoded blocks in $Q_r$.
    Only a consecutive ready prefix may be patched, compressed, appended and
    retired. Solid arrows show data/order dependencies; dashed arrows show
    fallback or decision updates. Boundary probing entails extra decoding.""",
         notes="""CODE BASIS
swbam-app/algorithms/markdup/swbam_markdup_mpi.cpp:
  MdBuildOnePassBoundaryPlan: probes min(span_count, 32) blocks at each end
  (overlapping probes are deduplicated). It builds rank coordinate boundaries,
  checks coverage by MPI_Allreduce(MIN), selects boundary candidates, routes
  them by coordinate owner, resolves duplicates and returns duplicate orders.
  MdAppendBoundaryCandidates / MdRouteBoundaryDuplicateOrders: metadata and
  QNAME information are exchanged, not complete decoded BAM records. H_r also
  identifies pre-resolved candidates. Empty ranks can trigger fallback.
  Probe extents are schematic and need not be exactly b +/- L. The top panel
  illustrates adjacent ranks on one reference, not every reference/empty-rank
  edge case. Owner routing is not restricted to only immediate neighbors.
  MdOnePassCaptureBatch: retains decoded storage; initializes provisional KEEP
  and applies cached boundary decisions. The main path suppresses candidates
  already resolved by the boundary plan. Window/local actions and routed remote
  decisions update per-record state and each block's unresolved count.
  MdOnePassPrepareOutput stops at the first unresolved block, patches flags,
  and builds output payload views; oversized legal payloads can be repacked.
  MdOnePassFinishOutput completes compression and flushes the body sink before
  retiring blocks; backing storage can be recycled at batch granularity.
swbam-app/algorithms/markdup/swbam_markdup_stream_mpi.cpp:
  MdWindowExpired / MdWindowKeepKey implement coordinate expiration and
  conservative boundary retention. This queue is controlled by the memory
  budget, not a claim of a fixed number of blocks for all inputs.
  Pending/ready states refer to unresolved duplicate decisions, not BGZF CRC
  or decode completion. A ready block cannot overtake an earlier pending one.
  This figure depicts default marking, not record removal or tag clearing.
"""),
    dict(stem="swbam_distributed_reordering", build=reordering,
         title="Distributed sorting and collation built on SWBAM",
         caption=r"""Distributed reordering built on the common BAM read and
    write pipelines. (a) Sort orders local metadata, chooses coordinate
    splitters from samples, redistributes records by coordinate range, and
    merges ordered source segments. (b) Collate routes QNAME hash bins to owner
    ranks and merges records into contiguous QNAME groups. Its ordering is
    based on bin/hash and QNAME, not global QNAME lexicographic order. Toy
    outputs illustrate the different ordering guarantees, not actual datasets
    or sampled splitters. (c) External-memory variants preserve their
    partitioning rules while using ordered runs, bounded exchanges and
    streaming merges. Algorithmic run files are distinct from output-body
    spool files.""",
         notes="""CODE BASIS
swbam-app/algorithms/sort/swbam_sort_optimized_mpi.cpp:
  OptimizedBamSortMPIImpl sorts metadata, makes local samples, chooses and
  broadcasts splitters, exchanges buckets, then k-way merges source segments.
  Stable tie information is retained in full keys; the figure's numeric toy
  uses distinct positions on one reference and omits those tie fields.
  External mode generates runs and streams owner-partitioned segments through
  bounded exchanges and merges. The bottom panel summarizes dataflow and does
  not replace its sample/splitter selection step.
swbam-app/algorithms/collate/swbam_collate_mpi.cpp:
  CollateOwner(bin,bins,ranks) = floor(bin*ranks/bins), clamped to the last
  rank, NOT bin % ranks. CollateCompare orders by bin, hash, QNAME, flag_order,
  then global_order. Full QNAME comparison resolves hash collisions.
  CollateExchangeMemory retains self-owned data locally and exchanges remote
  records in bounded ring chunks; received raw bytes go to their final arena.
  CollateMemoryLoserTree merges ordered segments into grouped output.
  CollateExtractRuns / CollateExchangeRuns / CollateMergeExternal use bounded
  arena and merge fan-in. The example qA/qB/qC ownership is illustrative;
  no specific hash value is asserted. Memory-mode diagrams use both metadata
  and raw record payloads, not candidate-only traffic as in markdup.
  Source-local order is preserved within sent segments. Both algorithms use
  common read/write pipelines; external temp runs and RankBodySink spool are
  separate storage mechanisms. I/O backend or spool policy is not prescribed.
"""),
]


def export_figure(spec):
    d.SCENE.clear()
    spec["build"]()
    d.validate_scene()
    svg = d.to_svg(spec["title"], spec["title"] + ". Academic method schematic.")
    stem = OUT / spec["stem"]
    stem.with_suffix(".svg").write_text(svg, encoding="utf-8")
    cairosvg.svg2png(bytestring=svg.encode(), output_width=4320,
                    write_to=str(stem.with_suffix(".png")))
    cairosvg.svg2pdf(bytestring=svg.encode(), write_to=str(stem.with_suffix(".pdf")))
    d.to_pptx(stem.with_suffix(".pptx"), spec["title"], COMMON_NOTES + "\n" + spec["notes"])
    tex = ("% Insert in the relevant Methods subsection; compile from overleaf/SWBAM/.\n"
           "\\begin{figure*}[!t]\n    \\centering\n"
           f"    \\includegraphics[width=\\textwidth]{{../figs/{spec['stem']}.pdf}}\n"
           f"    \\caption{{{spec['caption']}}}\n"
           f"    \\label{{fig:{spec['stem']}}}\n"
           "\\end{figure*}\n")
    stem.with_suffix(".tex").write_text(tex, encoding="utf-8")
    print("Exported:", spec["stem"])


def combined_deck():
    deck = Presentation()
    deck.slide_width, deck.slide_height = Inches(d.WIDTH / 100), Inches(d.HEIGHT / 100)
    deck.core_properties.title = "SWBAM Methods figures"
    deck.core_properties.author = "SWBAM"
    # Source presentations contain only native shapes; no relationship-bearing
    # images/charts are copied. Preserve any user edits to the existing overview.
    stems = ["swbam_overview", "swbam_batch_overlap",
             "swbam_distributed_reordering", "swbam_markdup_boundary"]
    for stem in stems:
        source = Presentation(OUT / (stem + ".pptx"))
        assert source.slide_width == deck.slide_width and source.slide_height == deck.slide_height
        original = source.slides[0]
        slide = deck.slides.add_slide(deck.slide_layouts[6])
        slide.background.fill.solid()
        slide.background.fill.fore_color.rgb = RGBColor(255, 255, 255)
        for shape in original.shapes:
            assert not shape._element.xpath(".//@r:embed | .//@r:link | .//@r:id")
            slide.shapes._spTree.insert_element_before(deepcopy(shape._element), "p:extLst")
        slide.notes_slide.notes_text_frame.text = original.notes_slide.notes_text_frame.text
    path = OUT / "swbam_method_figures.pptx"
    deck.save(path)
    assert len(Presentation(path).slides) == 4
    print("Combined editable deck:", path)


def main():
    for spec in FIGURES:
        export_figure(spec)
    combined_deck()


if __name__ == "__main__":
    main()
