#!/usr/bin/env python3
"""Polish visible wording on slides 1, 2 and 5 of the author's figure deck.

Run using the figure environment with python-pptx, Pillow and cairosvg.
The input, geometry, drawings, other slides and package assets are preserved.
"""

from hashlib import sha256
from pathlib import Path
import sys
from zipfile import ZipFile

sys.dont_write_bytecode = True

from pptx import Presentation

from swbam_sort_collate_revision import descendants, replace_text, validate_text


OUT = Path(__file__).resolve().parent
SOURCE = OUT.parent / "tmp" / "SWBAM_figures.pptx"
DESTINATION = OUT / "SWBAM_figures_polished.pptx"

# Keys are stable shape IDs in the author-edited deck. Font sizes use the
# figure generators' 1600 x 900 coordinate system (one unit = 0.6 pt).
CHANGES = {
    1: {
        2: ("(a) SWBAM library: data access and batch processing", None),
        4: ("Data access", None),
        5: ("BAM input backends\nMemory / POSIX / MPI-IO", 17),
        25: ("Rank-local BGZF blocks", 18),
        27: ("Double buffering", 20),
        34: ("Up to 64 blocks per batch", 17),
        36: ("Rank-local output storage", 20),
        47: ("Memory / temporary file", 18),
        52: ("Distributed\nBAM output\nMemory / MPI-IO", 17),
        54: ("Within each MPI rank: 1 MPE + 64 CPEs", 24),
        58: ("Prepare batch", None),
        61: ("Launch / join CPEs", 20),
        64: ("Batch post-processing", None),
        66: ("Batch-level control", 18),
        137: ("CPE processing components", None),
        138: ("Selected by the pipeline operator", 18),
        139: ("BAM read", None),
        141: ("Decompress", 18),
        145: ("Parse records", 18),
        149: ("Record\noperation", 18),
        151: ("BAM write", None),
        153: ("Serialize", None),
        157: ("Compress", None),
        159: ("SAM read", None),
        161: ("Text chunks", None),
        165: ("Parse SAM", 19),
        167: ("SAM write", None),
        169: ("BAM records", None),
        173: ("Format SAM", 19),
        176: ("Input batches", 18),
        178: ("Output batches", 18),
        179: ("(b) Batch-level pipeline composition", 26),
        180: ("Record counting", None),
        182: ("BAM read", None),
        186: ("Combine counts", 19),
        190: ("MPI reduction", 19),
        192: ("BAM filtering and rewriting", 21),
        194: ("BAM read + filter", None),
        198: ("Plan record packing", 19),
        202: ("BAM write", None),
        204: ("BAM \u2192 SAM", None),
        206: ("BAM read", None),
        210: ("Pass record batch", 19),
        214: ("SAM write", None),
        216: ("SAM \u2192 BAM", None),
        218: ("SAM read", None),
        222: ("Plan record packing", 19),
        226: ("BAM write", None),
        228: ("Sort, Collate, Fixmate and Markdup add ordering, grouping, "
              "cross-record state and MPI communication.", 20),
    },
    2: {
        2: ("(a) BAM read: overlap input reads with CPE processing", 26),
        4: ("Double buffering", 24),
        6: ("A: batch i", None),
        9: ("B: batch i+1", None),
        11: ("Compressed input buffers\nAlternate between A and B", 17),
        13: ("Time (schematic)", None),
        19: ("Prepare\ni", 17),
        22: ("Read batch i+1\ninto B", 18),
        25: ("Wait", None),
        28: ("Post-process\nbatch i", 18),
        31: ("Prepare\ni+1", 17),
        34: ("Read batch i+2\ninto A", 18),
        37: ("Wait", None),
        40: ("Post-process\nbatch i+1", 18),
        43: ("Decompress + parse + process\nBatch i (buffer A)", 18),
        46: ("Launch", 17),
        48: ("Join / validate", 17),
        50: ("Decompress + parse + process\nBatch i+1 (buffer B)", 18),
        53: ("Launch", 17),
        55: ("Join / validate", 17),
        59: ("Input read / CPE overlap", 18),
        60: ("Post-processing follows CPE completion and validation.", 19),
        61: ("(b) BAM write: overlap output writes with CPE processing", 26),
        63: ("Double buffering", 24),
        65: ("A: batch i", None),
        68: ("B: batch i\u22121", None),
        70: ("Compressed output buffers\nAlternate between A and B", 17),
        72: ("Time (schematic)", None),
        78: ("Prepare\ni", 17),
        81: ("Write batch i\u22121\nfrom B", 18),
        84: ("Read ahead\n(optional)", 18),
        87: ("Join / validate\nbatch i", 18),
        90: ("Prepare\ni+1", 17),
        93: ("Write batch i\nfrom A", 18),
        96: ("Read ahead\n(optional)", 18),
        99: ("Join / validate\nbatch i+1", 18),
        102: ("Serialize + compress\nBatch i (buffer A)", 20),
        105: ("Launch", 17),
        107: ("Complete", 17),
        109: ("Serialize + compress\nBatch i+1 (buffer B)", 20),
        112: ("Launch", 17),
        114: ("Complete", 17),
        118: ("Output write / CPE overlap", 18),
        119: ("Packed BGZF payloads require compression only.", 19),
        120: ("Read and write schedules are shown separately; "
              "each core group runs one CPE kernel at a time.", 19),
    },
    5: {
        2: ("(a) Resolve boundary duplicates before streaming", 26),
        4: ("Probe the first and last BGZF blocks of each rank", 23),
        5: ("Rank r\u22121: last blocks", 21),
        6: ("Rank r: first blocks", 21),
        42: ("Boundary region", 18),
        43: ("CPE decompression and candidate extraction\n"
             "Exchange candidate metadata, not full BAM records", 18),
        46: ("Full coverage\non all ranks?", 18),
        48: ("Yes", None),
        50: ("Coordinate owner rank", 23),
        52: ("r\u22121: key k  |  score 70", None),
        55: ("r: key k  |  score 60", None),
        57: ("Resolve matching candidates", 18),
        59: ("Return", None),
        61: ("Boundary decisions H\u1d63", 22),
        63: ("r\u22121   :   KEEP", None),
        66: ("r      :   DUP", None),
        68: ("Indexed by source record ID", 18),
        71: ("No", None),
        72: ("Retain boundary candidates\nin the conservative window", 19),
        73: ("(b) Mark duplicates and write finalized blocks in order", 26),
        75: ("Decode + extract (CPE)", 21),
        76: ("Decoded blocks B\nCandidates C", 20),
        78: ("Local C", None),
        80: ("Candidate window W\u1d63", None),
        81: ("Local and received candidates", 17),
        83: ("Pair keys", 19),
        86: ("Single-read keys", 18),
        88: ("Exclude candidates resolved by H\u1d63", 18),
        90: ("Remote owner ranks", 21),
        91: ("Resolve duplicate keys\nReturn decisions", 18),
        95: ("Remote candidates", None),
        97: ("Returned decisions", None),
        99: ("Apply boundary decisions\nH\u1d63: KEEP / DUP", 21),
        101: ("Retain decoded blocks\n(provisional KEEP)", None),
        104: ("Pending queue Q\u1d63 (input order)", None),
        105: ("Tail", None),
        108: ("B\u2080   Finalized", 22),
        112: ("B\u2081   Finalized", 22),
        116: ("B\u2082   Pending", None),
        120: ("B\u2083   Finalized, waiting", 21),
        125: ("Local decisions", None),
        130: ("Write finalized blocks in input order", 18),
        131: ("Stop at the first unresolved block", None),
        134: ("Update duplicate flags", 19),
        137: ("CPE BGZF compression", None),
        140: ("Append to rank output", 19),
        143: ("Release block storage", 19),
        148: ("Boundary probes require\nlimited extra decoding.", 17),
    },
}


def geometry(slide):
    return [(s.shape_id, s.left, s.top, s.width, s.height, s.rotation)
            for s in descendants(slide.shapes)]


def main():
    source_hash = sha256(SOURCE.read_bytes()).hexdigest()
    prs = Presentation(SOURCE)
    assert len(prs.slides) == 5
    replacements = {}
    for number, edits in CHANGES.items():
        slide = prs.slides[number - 1]
        before = geometry(slide)
        shapes = {s.shape_id: s for s in descendants(slide.shapes)}
        changed = 0
        for sid, (text, size) in edits.items():
            shape = shapes[sid]
            assert shape.has_text_frame and shape.text, (number, sid)
            if shape.text == text:
                continue
            replace_text(shape, text, size=size)
            changed += 1
        assert geometry(slide) == before
        validate_text(slide)
        replacements[str(slide.part.partname).lstrip("/")] = slide.part.blob
        print(f"Slide {number}: {changed} text boxes revised; geometry unchanged")

    # Replace only the three slide XML parts, not the whole python-pptx package.
    # In particular, slides 3/4, notes, themes and all drawing relationships
    # remain byte-for-byte identical to the author's file.
    with ZipFile(SOURCE) as original, ZipFile(DESTINATION, "w") as result:
        for entry in original.infolist():
            result.writestr(entry, replacements.get(
                entry.filename, original.read(entry.filename)))
    with ZipFile(SOURCE) as original, ZipFile(DESTINATION) as result:
        assert result.testzip() is None
        assert original.namelist() == result.namelist()
        changed_parts = [name for name in original.namelist()
                         if original.read(name) != result.read(name)]
        assert set(changed_parts) == set(replacements), changed_parts
    assert sha256(SOURCE.read_bytes()).hexdigest() == source_hash
    assert len(Presentation(DESTINATION).slides) == 5
    print("Changed parts:", ", ".join(changed_parts))
    print(DESTINATION)


if __name__ == "__main__":
    main()
