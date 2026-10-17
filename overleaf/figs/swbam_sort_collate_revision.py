#!/usr/bin/env python3
"""Revise slide 3 of the author's edited deck, retaining the other slides.

Run with python-pptx, Pillow and cairosvg installed. All figure objects remain
native PowerPoint shapes. The input in overleaf/tmp is never overwritten.
"""

from copy import deepcopy
from pathlib import Path
import sys
from zipfile import ZipFile

sys.dont_write_bytecode = True

import cairosvg
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE_TYPE
from pptx.enum.text import MSO_ANCHOR, MSO_AUTO_SIZE, PP_ALIGN
from pptx.util import Inches, Pt

from swbam_method_redesign import (
    Canvas, W, H, INK, TEXT, GRAY, BLUE, GREEN, YELLOW, GOLD, PURPLE, WHITE,
)


OUT = Path(__file__).resolve().parent
SOURCE = OUT.parent / 'tmp' / 'swbam_method_figures_redesigned_v2.pptx'
DECK = OUT / 'swbam_method_figures_redesigned_v3.pptx'
SINGLE = OUT / 'swbam_sort_collate_reordering_v3.pptx'
PREVIEW = Path('/tmp/swbam-redesign/sort-collate-v3')


def descendants(shapes):
    for shape in shapes:
        yield shape
        if shape.shape_type == MSO_SHAPE_TYPE.GROUP:
            yield from descendants(shape.shapes)


def replace_text(shape, text, size=None, rect=None, bold=None, align=None):
    """Retain the author's font/color while updating text and its fixed box."""
    runs = [r for p in shape.text_frame.paragraphs for r in p.runs]
    first = runs[0] if runs else None
    alignment = shape.text_frame.paragraphs[0].alignment or PP_ALIGN.CENTER
    if align is not None:
        alignment = PP_ALIGN.LEFT if align == 'left' else PP_ALIGN.CENTER
    font_size = size if size is not None else first.font.size.pt / 0.6
    is_bold = bold if bold is not None else bool(first.font.bold)
    color = TEXT
    if first is not None and first.font.color.type is not None:
        try:
            color = str(first.font.color.rgb)
        except AttributeError:
            pass
    if rect is not None:
        shape.left, shape.top, shape.width, shape.height = [Inches(v / 120) for v in rect]
    tf = shape.text_frame
    tf.clear()
    tf.word_wrap = False
    tf.auto_size = MSO_AUTO_SIZE.NONE
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    for i, line in enumerate(text.split('\n')):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = alignment
        p.space_before = p.space_after = Pt(0)
        p.line_spacing = Pt(font_size * 0.6 * 1.12)
        r = p.add_run()
        r.text = line
        r.font.name = 'Arial'
        r.font.size = Pt(font_size * 0.6)
        r.font.bold = is_bold
        r.font.italic = False
        r.font.color.rgb = RGBColor.from_string(color)


def storage_node(c, name, x, title, temporary_label):
    c.shape(name + ' / frame', x, 603, 170, 164, WHITE, INK, width=1.5)
    c.text(name + ' / heading', x + 8, 608, 154, 42,
           title, 19, True)
    c.shape(name + ' / memory choice', x + 8, 654, 154, 42,
            BLUE, INK, width=1)
    c.text(name + ' / RAM label', x + 14, 662, 61, 24, 'RAM', 18)
    for j in range(3):
        c.shape(name + f' / RAM cell {j}', x + 84 + j * 21, 666,
                21, 18, WHITE, INK, width=1)
    c.shape(name + ' / external choice', x + 8, 707, 154, 50,
            PURPLE, INK, width=1)
    c.text(name + ' / temporary label', x + 13, 712, 99, 39,
           'Temporary\n' + temporary_label, 17)
    for j in range(2, -1, -1):
        c.shape(name + f' / temporary extent {j}',
                x + 124 + j * 3, 720 - j * 3, 24, 29,
                WHITE if j else PURPLE, INK, kind='paper', width=1)
    for j in range(2):
        c.line(name + f' / record line {j}',
               [(x + 129, 729 + j * 6), (x + 143, 729 + j * 6)], INK, 0.9)


def storage_panel():
    c = Canvas('Shared workflow with alternative storage',
        'The six logical stages are shared by in-memory and external-memory '
        'execution. Each storage box contains two alternatives selected by '
        'the algorithm mode; its RAM and temporary-file rows are not serial '
        'stages. External execution sorts successive bounded input chunks, '
        'stores local runs, exchanges bounded chunks into received segments, '
        'and consolidates segments under a memory-derived merge fan-in F. '
        'Only the final merge enters BGZF output. The dashed left arrow is '
        'an iteration over new input, not a readback of runs to re-sort them. '
        'Temporary extents contain metadata and raw BAM bytes; each icon '
        'does not imply a separate file. Some runs may remain resident. '
        'Algorithm temporary storage is independent of final rank-body spool.')
    c.panel('Unified storage panel', 664, 512, 908, 373,
            '(c) Shared workflow with alternative storage', WHITE)

    c.shape('Memory legend swatch', 690, 565, 18, 18, BLUE, INK, width=1)
    c.text('Memory legend', 716, 559, 118, 29, 'In-memory', 18, align='left')
    c.shape('External legend swatch', 859, 565, 18, 18, PURPLE, INK, width=1)
    c.text('External legend', 885, 559, 109, 29, 'External', 18, align='left')
    c.line('External iteration legend line', [(1050,574), (1104,574)], GOLD, 1.7,
           dash=True, head=True)
    c.text('External iteration legend label', 1116, 559, 421, 29,
           'External mode only', 18, color=GOLD, align='left')

    c.box('Local sort operation', 686, 658, 105, 76,
          'Local\nsort', YELLOW, 20)
    storage_node(c, 'Local ordered storage', 813, 'Local\nsorted data', 'runs')
    c.box('Chunked MPI operation', 1005, 658, 110, 76,
          'Chunked\nexchange', WHITE, 18)
    storage_node(c, 'Received segment storage', 1137, 'Received\nsegments', 'segments')
    c.box('Multiway merge operation', 1329, 658, 105, 76,
          'Multi-way\nmerge', YELLOW, 18)
    c.box('BGZF output operation', 1456, 658, 94, 76,
          'BGZF\noutput', GREEN, 19)
    for name, x1, x2 in [
            ('Save local ordered data', 791, 813),
            ('Load data for exchange', 983, 1005),
            ('Save received ordered data', 1115, 1137),
            ('Read ordered segments', 1307, 1329),
            ('Write final merge', 1434, 1456)]:
        c.arrow(name, x1, 700, x2, 700, width=1.8)
    c.text('Final merge pass', 1433, 623, 117, 24,
           'Final pass', 17, color=GRAY)

    c.line('External run generation iterations',
           [(898,767), (898,789), (738,789), (738,737)],
           GOLD, 1.7, head=True, dash=True)
    c.text('External input iteration label', 686, 800, 334, 26,
           'Repeat for each input chunk', 17, color=GOLD)
    c.line('External intermediate merge iterations',
           [(1381,734), (1381,789), (1222,789), (1222,770)],
           GOLD, 1.7, head=True, dash=True)
    c.text('External merge iteration label', 1113, 800, 433, 26,
           'Additional passes if segments exceed F', 17, color=GOLD)
    c.text('Temporary data representation', 686, 846, 368, 25,
           'Stored data: metadata + raw BAM bytes', 17, color=GRAY, align='left')
    c.text('Fan-in explanation', 1113, 839, 433, 39,
           'F: merge fan-in\nDetermined by the memory budget', 17, color=GRAY)
    c.validate()
    return c


def revise_text(slide):
    lookup = {s.shape_id: s for s in descendants(slide.shapes)}
    changes = {
        2: ('(a) Distributed sorting and collation (six MPI ranks)', 26, None),
        5: ('BAM read pipeline', 22, None),
        6: ('Decode + extract (MPE / CPE)', 16, None),
        8: ('Local sorting', 22, None),
        11: ('Redistribution', 21, None),
        12: ('By destination rank', 16, None),
        14: ('Sorted segments', 22, None),
        15: ('Grouped by source rank', 17, None),
        17: ('Merge + BAM write pipeline', 21, None),
        439: ('Chunked\nMPI exchange', 22, (646, 199, 173, 62)),
        444: ('Local data:\nno network transfer', 17, (646, 345, 178, 49)),
        445: ('Destination rank', 18, (350, 451, 167, 28)),
        446: ('Segments by source', 18, (859, 425, 286, 29)),
        447: ('Merged records', 18, (1212, 425, 198, 29)),
        449: ('(b) Partitioning rules', 24, None),
        451: ('Samples \u2192 5 splitters \u2192 6 ranges', 19, (165, 570, 454, 28)),
        483: ('QNAME hash \u2192 66 bins \u2192 6 ranks', 19, (177, 729, 444, 28)),
        566: ('Merge', 18, None),
        571: ('Decoded BAM records', 18, (140, 425, 198, 29)),
        572: ('Sorted metadata', 18, (425, 425, 184, 29)),
        573: ('Records with the same QNAME are contiguous', 19, (49, 840, 577, 30)),
        578: ('Destination rank', 18, (66, 451, 167, 28)),
        579: ('BGZF blocks', 18, (1420, 425, 140, 29)),
    }
    for sid, (text, size, rect) in changes.items():
        replace_text(lookup[sid], text, size, rect)
    for r in range(6):
        for prefix in ('Coordinate owner range', 'Collate owner range'):
            name = f'{prefix} {r} / label'
            shape = next(s for s in lookup.values() if s.name == name)
            replace_text(shape, f'rank {r}', size=18)


def append_native_elements(target, source):
    next_id = max(s.shape_id for s in descendants(target.shapes)) + 1
    inserted = []
    for source_shape in source.shapes:
        node = deepcopy(source_shape._element)
        for prop in node.xpath('.//p:cNvPr'):
            prop.set('id', str(next_id))
            next_id += 1
        target.shapes._spTree.insert_element_before(node, 'p:extLst')
        inserted.append(target.shapes[-1])
    return inserted


def validate_text(slide):
    c = Canvas('Text layout checks', '')
    ids = []
    for shape in descendants(slide.shapes):
        ids.append(shape.shape_id)
        assert shape.shape_type != MSO_SHAPE_TYPE.PICTURE, shape.name
        if not shape.has_text_frame or not shape.text:
            continue
        r = shape.text_frame.paragraphs[0].runs[0]
        x, y, w, h = [v / 914400 * 120 for v in
                       (shape.left, shape.top, shape.width, shape.height)]
        c.text(shape.name, x, y, w, h, shape.text, r.font.size.pt / 0.6,
               bool(r.font.bold))
    assert len(ids) == len(set(ids)), 'Duplicate PowerPoint shape IDs'
    c.validate()


def preserve_package(source, output, slide):
    # Only slide 3 and its notes change. This also preserves the author's
    # other slides, theme, metadata, and embedded assets byte for byte.
    replacements = {
        str(slide.part.partname).lstrip('/'): slide.part.blob,
        str(slide.notes_slide.part.partname).lstrip('/'): slide.notes_slide.part.blob,
    }
    with ZipFile(source) as original, ZipFile(output, 'w') as result:
        for entry in original.infolist():
            result.writestr(entry, replacements.get(entry.filename, original.read(entry.filename)))
    with ZipFile(source) as original, ZipFile(output) as result:
        assert result.testzip() is None
        assert set(original.namelist()) == set(result.namelist())
        for name in original.namelist():
            if name not in replacements:
                assert original.read(name) == result.read(name), name


def main():
    PREVIEW.mkdir(parents=True, exist_ok=True)
    prs = Presentation(SOURCE)
    assert len(prs.slides) == 4
    slide = prs.slides[2]
    assert slide.shapes[0].shape_id == 2
    revise_text(slide)
    # Remove old panel c, two crossing callout lines, and an empty off-slide box.
    for shape in list(slide.shapes):
        if 509 <= shape.shape_id <= 565 or shape.shape_id in (567, 574, 580):
            slide.shapes._spTree.remove(shape._element)
    panel = storage_panel()
    temporary = Presentation()
    temporary.slide_width, temporary.slide_height = Inches(W / 120), Inches(H / 120)
    drawn = panel.add_slide(temporary)
    append_native_elements(slide, drawn)

    # Keep the author's partition callout, routed through the column gap.
    callout = Canvas('Partition rule callout', '')
    callout.arrow('Local sorting to partition rules', 620, 415, 620, 512,
                  INK, dash=True, width=1.6)
    append_native_elements(slide, callout.add_slide(temporary))
    slide.notes_slide.notes_text_frame.text = (
        'Distributed sorting and collation: revised author-edited slide 3.\n\n'
        'Panel (a): cell colors and numbers denote eventual destination ranks, '
        'including before Sort splitters are selected. The extracted data are '
        'decoded BAM bytes plus metadata. Local sorting rearranges metadata. '
        'The same-QNAME grouping performed by Collate is not globally '
        'lexicographic QNAME sorting. Rank-local output is therefore called '
        'merged records rather than globally sorted data. Local data avoid '
        'network transfer but may incur local copies. src0-src5 identify '
        'source ranks. The six segment glyphs summarize source contributions: '
        'sources can be empty or supply multiple segments in external mode. '
        'Merged records may interleave sources. Rank bodies are joined under '
        'one BAM header with one final BGZF EOF.\n\n'
        'Panel (b): Sort gathers local samples, selects five global splitters '
        'for six ranks, and broadcasts them. Collate uses '
        'owner=floor(bin*nranks/nbins); 66 bins are an example. Its comparator '
        'uses bin, hash, full QNAME, group order and global input order.\n\n'
        + panel.notes + '\n\n'
        'Code: swbam-app/algorithms/sort/swbam_sort_optimized_mpi.cpp; '
        'swbam-app/algorithms/collate/swbam_collate_mpi.cpp. '
        'Terminology follows standard local sorting, partitioning, data '
        'exchange, sorted runs, and multi-way merging, as used in the '
        'reference RabbitBAM/RabbitBAM-SORT paper. Visual style follows '
        'overleaf/reference/plot.pptx. The diagram summarizes separate '
        'Sort/Collate implementations; it does not claim identical code paths.'
    )
    validate_text(slide)
    preserve_package(SOURCE, DECK, slide)

    single = Presentation()
    single.slide_width, single.slide_height = prs.slide_width, prs.slide_height
    single.core_properties.title = 'SWBAM distributed sorting and collation'
    single.core_properties.author = 'SWBAM'
    only = single.slides.add_slide(single.slide_layouts[6])
    for shape in slide.shapes:
        only.shapes._spTree.insert_element_before(deepcopy(shape._element), 'p:extLst')
    only.notes_slide.notes_text_frame.text = slide.notes_slide.notes_text_frame.text
    single.save(SINGLE)

    svg = panel.svg()
    (PREVIEW / 'panel_c.svg').write_text(svg, encoding='utf-8')
    # SVG viewBox crops the vector panel without editing a rendered bitmap.
    cropped = svg.replace('width="1600" height="900" viewBox="0 0 1600 900"',
                          'width="930" height="393" viewBox="653 502 930 393"', 1)
    cairosvg.svg2png(bytestring=cropped.encode(),
                    write_to=str(OUT / 'swbam_sort_collate_panel_c_v3.png'),
                    output_width=1860, output_height=786)
    print(DECK)
    print(SINGLE)
    print('Verified: only slide 3 and its notes differ in the complete deck.')


if __name__ == '__main__':
    main()
